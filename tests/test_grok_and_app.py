import json
from io import BytesIO

import httpx
import pytest
from docx import Document
from fastapi.testclient import TestClient

from interpreter_summary.config import Settings
from interpreter_summary.export import markdown_to_docx
from interpreter_summary.grok import GrokClient, extract_output_text
from interpreter_summary.samples import SAMPLE_SUMMARY_MARKDOWN, write_sample_pdf
from interpreter_summary.service import summarize_pdf
from interpreter_summary.web.app import app


WRONG_PAGE_MARKDOWN = """# Interpreting Interpreter: Tokens, Not Tonnage

## The Takeaway

Scholar argues that the copper is a covenant token.

## The Q&A

### What is on the list?

The fragment lists ingots of bright copper.

## The Summary

The fragment lists ingots (link to "The warehouse fragment"; page 99).

## Video Script

| # | Text | Image |
| --- | --- | --- |
| 1 | Ten ingots, right after an oath. | Bright copper |
| 10 | Check out the full article, and I'll see you next time. | Title page |
"""


PARAPHRASE_QUOTE_MARKDOWN = """# Interpreting Interpreter: Tokens, Not Tonnage

## The Takeaway

Scholar argues that the copper is a covenant token.

## The Q&A

### What is Fragment W?

A short list of bright copper ingots after an oath formula.

## The Summary

The fragment lists ingots (link to "The warehouse fragment"; page 1).

As Scholar concludes (link to "although copper can be cargo"; page 3):

> although copper can be cargo, in this fragment it obviously memorializes a promise.

## The Reflection

I like the caution.

## Video Script

| # | Text | Image |
| --- | --- | --- |
| 1 | Ten ingots, right after an oath. | Bright copper |
| 2 | Check out the full article, and I'll see you next time. | Title page |
"""


def _handler(markdown: str, seen: list | None = None):
    def respond(request: httpx.Request) -> httpx.Response:
        if request.method == "POST" and request.url.path == "/v1/files":
            return httpx.Response(200, json={"id": "file-123", "filename": "article.pdf"})
        if request.method == "POST" and request.url.path == "/v1/responses":
            payload = json.loads(request.content.decode() or "{}")
            if seen is not None:
                seen.append(payload)
            if payload.get("stream"):
                event = json.dumps({"type": "response.output_text.delta", "delta": markdown})
                body = f"data: {event}\n\ndata: [DONE]\n\n"
                return httpx.Response(200, text=body)
            return httpx.Response(200, json={"output_text": markdown})
        if request.method == "DELETE" and request.url.path.startswith("/v1/files/"):
            return httpx.Response(200, json={"deleted": True})
        return httpx.Response(404, json={"error": "unhandled"})

    return respond


def _file_attachment(payload: dict) -> dict | None:
    for item in payload.get("input") or []:
        content = item.get("content")
        if not isinstance(content, list):
            continue
        for part in content:
            if isinstance(part, dict) and part.get("type") == "input_file":
                return part
    return None


@pytest.mark.asyncio
async def test_summarize_pdf_with_mocked_grok(tmp_path):
    pdf_path = write_sample_pdf(tmp_path / "article.pdf")
    settings = Settings(xai_api_key="test-key", xai_model="grok-4.6")
    seen: list[dict] = []
    transport = httpx.MockTransport(_handler(SAMPLE_SUMMARY_MARKDOWN, seen))
    async with httpx.AsyncClient(transport=transport, base_url="https://api.x.ai") as raw:
        client = GrokClient(settings, client=raw)
        result = await summarize_pdf(
            pdf_path.read_bytes(),
            "article.pdf",
            settings,
            client=client,
        )
    assert result.title == "Interpreting Interpreter: Tokens, Not Tonnage"
    assert result.locator_count == 7
    assert result.locators_corrected == 0
    assert result.locators_dropped == 0
    assert result.quotes_kept == 1
    assert result.quotes_corrected == 0
    assert result.quotes_dropped == 0
    assert result.page_count == 3
    assert result.journal_page_start == 1
    assert result.summary_word_count > 0
    assert "target" in result.summary_length_note
    assert "covenant token" in result.takeaway
    assert "Fragment W" in result.qa
    assert "In this article" in result.summary
    assert result.video_row_count == 10
    assert "see you next time" in result.video_script
    assert "Reflection" not in result.markdown
    assert "3 questions" in result.qa_length_note
    assert "video 10 rows" in result.video_length_note
    document = Document(BytesIO(markdown_to_docx(result.markdown)))
    texts = [paragraph.text for paragraph in document.paragraphs]
    assert "The Takeaway" in texts
    assert "The Q&A" in texts
    assert "The Summary" in texts
    assert "Video Script" in texts
    assert "The Reflection" not in texts
    assert len(document.tables) == 1
    assert [cell.text for cell in document.tables[0].rows[0].cells] == ["#", "Text", "Image"]
    assert "see you next time" in document.tables[0].rows[-1].cells[1].text
    assert seen and seen[0]["stream"] is True
    attachment = _file_attachment(seen[0])
    assert attachment is not None
    assert attachment["file_id"] == "file-123"
    assert "tool_choice" not in seen[0]


@pytest.mark.asyncio
async def test_summarize_pdf_strips_preamble_glued_to_the_title(tmp_path):
    pdf_path = write_sample_pdf(tmp_path / "article.pdf")
    settings = Settings(xai_api_key="test-key", xai_model="grok-4.6")
    chatter = (
        "I'll pull exact paragraph openings, printed pages, and the closing quotation."
    )
    transport = httpx.MockTransport(_handler(chatter + SAMPLE_SUMMARY_MARKDOWN))
    async with httpx.AsyncClient(transport=transport, base_url="https://api.x.ai") as raw:
        client = GrokClient(settings, client=raw)
        result = await summarize_pdf(
            pdf_path.read_bytes(),
            "article.pdf",
            settings,
            client=client,
        )
    assert result.title == "Interpreting Interpreter: Tokens, Not Tonnage"
    assert "I'll pull" not in result.markdown
    assert result.quotes_kept == 1
    document = Document(BytesIO(markdown_to_docx(result.markdown)))
    assert "I'll pull" not in document.paragraphs[0].text
    assert document.paragraphs[0].text.startswith("Interpreting Interpreter: Tokens, Not Tonnage")


@pytest.mark.asyncio
async def test_summarize_pdf_surfaces_a_length_note_outside_the_post(tmp_path):
    pdf_path = write_sample_pdf(tmp_path / "article.pdf")
    settings = Settings(xai_api_key="test-key", xai_model="grok-4.6")
    note = (
        "<!-- length-note: below the band because this is a close reading of one fragment -->\n"
    )
    transport = httpx.MockTransport(_handler(SAMPLE_SUMMARY_MARKDOWN + note))
    async with httpx.AsyncClient(transport=transport, base_url="https://api.x.ai") as raw:
        client = GrokClient(settings, client=raw)
        result = await summarize_pdf(
            pdf_path.read_bytes(),
            "article.pdf",
            settings,
            client=client,
        )
    assert "length-note" not in result.markdown
    assert "close reading" not in result.markdown
    assert "close reading" in result.summary_length_note
    document = Document(BytesIO(markdown_to_docx(result.markdown)))
    assert all("close reading" not in paragraph.text for paragraph in document.paragraphs)
    assert "closing quote" in result.closing_quote_length_note
    assert "30–45" in result.takeaway_length_note or "30-45" in result.takeaway_length_note


@pytest.mark.asyncio
async def test_summarize_pdf_corrects_wrong_locator_pages(tmp_path):
    pdf_path = write_sample_pdf(tmp_path / "article.pdf")
    settings = Settings(xai_api_key="test-key", xai_model="grok-4.6")
    transport = httpx.MockTransport(_handler(WRONG_PAGE_MARKDOWN))
    async with httpx.AsyncClient(transport=transport, base_url="https://api.x.ai") as raw:
        client = GrokClient(settings, client=raw)
        result = await summarize_pdf(
            pdf_path.read_bytes(),
            "article.pdf",
            settings,
            client=client,
        )
    assert result.locator_count == 1
    assert result.locators_corrected == 1
    assert result.locators_dropped == 0
    assert '(link to "The warehouse fragment"; page 1)' in result.markdown
    assert "page 99" not in result.markdown
    assert "The Q&A" in result.markdown


@pytest.mark.asyncio
async def test_summarize_pdf_repairs_paraphrased_block_quote(tmp_path):
    pdf_path = write_sample_pdf(tmp_path / "article.pdf")
    settings = Settings(xai_api_key="test-key", xai_model="grok-4.6")
    transport = httpx.MockTransport(_handler(PARAPHRASE_QUOTE_MARKDOWN))
    async with httpx.AsyncClient(transport=transport, base_url="https://api.x.ai") as raw:
        client = GrokClient(settings, client=raw)
        result = await summarize_pdf(
            pdf_path.read_bytes(),
            "article.pdf",
            settings,
            client=client,
        )
    assert result.quotes_corrected == 1
    assert result.quotes_dropped == 0
    assert "obviously" not in result.markdown
    assert "more plausibly memorializes a promise" in result.markdown
    assert "The Reflection" not in result.markdown
    assert "I like the caution" not in result.markdown
    assert "Video Script" in result.markdown


def test_extract_output_text_from_choices():
    payload = {"choices": [{"message": {"content": "hello from chat completions"}}]}
    assert extract_output_text(payload) == "hello from chat completions"


def test_health_and_index():
    client = TestClient(app)
    health = client.get("/health")
    assert health.status_code == 200
    assert health.json()["status"] == "ok"
    home = client.get("/")
    assert home.status_code == 200
    assert "Interpreting" in home.text
    assert "Q&amp;A" in home.text
    assert "video-script" in home.text
    assert "Reflection" not in home.text


def test_summarize_endpoint_rejects_non_pdf():
    client = TestClient(app)
    response = client.post(
        "/api/summarize",
        files={"pdf": ("notes.txt", b"hello", "text/plain")},
    )
    assert response.status_code == 400


def test_export_docx_roundtrip():
    client = TestClient(app)
    response = client.post("/api/export/docx", json={"markdown": SAMPLE_SUMMARY_MARKDOWN})
    assert response.status_code == 200
    assert response.content[:2] == b"PK"
    document = Document(BytesIO(response.content))
    assert document.tables
    assert "The Q&A" in [paragraph.text for paragraph in document.paragraphs]
    html = client.post(
        "/api/summarize",
        files={"pdf": ("notes.txt", b"hello", "text/plain")},
    )
    assert html.status_code == 400


def test_export_html_contains_video_table():
    import markdown as md

    html = md.markdown(SAMPLE_SUMMARY_MARKDOWN, extensions=["extra", "sane_lists", "nl2br"])
    assert "<table" in html
    assert "see you next time" in html
    assert "<h2>The Q&amp;A</h2>" in html or "<h2>The Q&A</h2>" in html


def test_extract_handles_output_text_field():
    assert extract_output_text({"output_text": "plain"}) == "plain"


@pytest.mark.asyncio
async def test_summarize_file_uses_completed_event_when_no_deltas_arrive():
    markdown = "# Interpreting Interpreter: Quiet Model\n\n## The Takeaway\n\nDone.\n"

    def respond(request: httpx.Request) -> httpx.Response:
        if request.url.path != "/v1/responses":
            return httpx.Response(404)
        event = json.dumps(
            {"type": "response.completed", "response": {"output_text": markdown}}
        )
        return httpx.Response(200, text=f"data: {event}\n\ndata: [DONE]\n\n")

    settings = Settings(xai_api_key="test-key", request_timeout_seconds=1200)
    transport = httpx.MockTransport(respond)
    async with httpx.AsyncClient(transport=transport, base_url="https://api.x.ai") as raw:
        client = GrokClient(settings, client=raw)
        text = await client.summarize_file(
            "file-123",
            system_prompt="system",
            user_prompt="user",
        )
    assert text.startswith("# Interpreting Interpreter: Quiet Model")
    assert settings.request_timeout_seconds == 1200
