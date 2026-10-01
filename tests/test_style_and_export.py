from io import BytesIO

from docx import Document

from interpreter_summary.export import (
    markdown_to_docx,
    parse_sections,
    strip_preamble,
    strip_reflection_section,
)
from interpreter_summary.length import (
    closing_quote_length_note,
    extract_length_note,
    length_hint,
    qa_length_note,
    question_count,
    summary_length_note,
    summary_prose,
    summary_word_target,
    takeaway_length_note,
    video_length_note,
    video_row_count,
)
from interpreter_summary.locators import locator_count, verify_locators
from interpreter_summary.pdf_utils import PdfCorpus, PdfPage, normalize_for_match, readable_text
from interpreter_summary.prompts import SYSTEM_PROMPT, build_user_prompt
from interpreter_summary.quotes import verify_quotes
from interpreter_summary.samples import SAMPLE_SUMMARY_MARKDOWN
from interpreter_summary.style import load_style_text


def test_system_prompt_requires_current_format():
    lowered = SYSTEM_PROMPT.lower()
    assert "The Takeaway" in SYSTEM_PROMPT
    assert "The Q&A" in SYSTEM_PROMPT
    assert "The Summary" in SYSTEM_PROMPT
    assert "Video Script" in SYSTEM_PROMPT
    assert "Do not write a Reflection" in SYSTEM_PROMPT
    assert "see you next time" in lowered
    assert "exactly three" in lowered
    assert "printed journal page" in lowered
    assert "viewer" in lowered
    assert "verbatim" in lowered
    assert "briefing" in lowered
    assert "cheerleader" in lowered
    assert "30–45" in SYSTEM_PROMPT or "30-45" in SYSTEM_PROMPT
    assert "35–70" in SYSTEM_PROMPT or "35-70" in SYSTEM_PROMPT
    assert "length-note" in SYSTEM_PROMPT
    assert "scales with article length" in lowered
    assert "briefly summarizing" in lowered
    assert "hagoth" in lowered
    assert "output only the post" in lowered
    assert "throat-clearing" in lowered
    assert "delve" in lowered
    assert "chatbot" in lowered
    assert "boilerplate" in lowered


def test_style_guide_teaches_current_voice():
    guide = load_style_text()
    lowered = guide.lower()
    assert "briefing" in lowered
    assert "cheerleader" in lowered
    assert "On Abstracting Thought" in guide
    assert "The Q&A" in guide
    assert "Video Script" in guide
    assert "see you next time" in lowered
    assert "Do not write a Reflection" in guide
    assert "In this article" in guide
    assert "200–330" in guide or "200-330" in guide
    assert "450–750" in guide or "450-750" in guide
    assert "30–45" in guide or "30-45" in guide
    assert "35–70" in guide or "35-70" in guide
    assert "60–110" in guide or "60-110" in guide
    assert "length-note" in lowered
    assert "bowen" in lowered
    assert "briefly summarizing" in lowered
    assert "hagoth is one of the coolest" in lowered
    assert "no preamble" in lowered
    assert "Clarity first" in guide
    assert "sheds light" in guide
    assert "Banned habits" in guide
    assert "green cacao" in lowered
    assert "plain" in lowered
    assert "Alma 63" in guide
    assert "Title page" in guide


def test_user_prompt_includes_style_and_asks_for_video_script():
    prompt = build_user_prompt(
        "House style goes here",
        extra_instructions="Keep it short.",
        pagination_hint="Printed journal pages are 425–450.",
        length_hint_text="The Summary: about 450–900 words.",
    )
    assert "House style goes here" in prompt
    assert "Video Script" in prompt
    assert "Do not write a Reflection" in prompt
    assert "Do not include a video script" not in prompt
    assert "Keep it short." in prompt
    assert "Printed journal pages are 425–450." in prompt
    assert "The Summary: about 450–900 words." in prompt


def _qa(answers: list[int], *, extra_question: bool = False) -> str:
    parts = []
    for index, words in enumerate(answers, start=1):
        parts.append(f"### Question {index}?\n\n" + " ".join(["word"] * words))
    if extra_question:
        parts.append("### And a fourth?\n\n" + " ".join(["word"] * 40))
    return "\n\n".join(parts)


def test_summary_word_targets_scale_with_pages_and_tighten_on_body_words():
    assert summary_word_target(8) == (200, 330, 400)
    assert summary_word_target(12) == (200, 330, 400)
    assert summary_word_target(16) == (330, 500, 650)
    assert summary_word_target(24) == (330, 500, 650)
    assert summary_word_target(26) == (380, 620, 800)
    assert summary_word_target(40) == (380, 620, 800)
    assert summary_word_target(42) == (450, 750, 950)
    assert summary_word_target(82) == (450, 750, 950)
    # Figure-heavy or note-heavy: body words select the shorter band.
    assert summary_word_target(60, 9236) == (380, 620, 800)
    assert summary_word_target(48, 7538) == (330, 500, 650)
    # Squire 70_04: 16 pages and ~5.8k body words stay on the page band.
    assert summary_word_target(16, 5780) == (330, 500, 650)
    # A denser text does not raise the page band.
    assert summary_word_target(24, 8488) == (330, 500, 650)
    # A thin extract is ignored.
    assert summary_word_target(30, 80) == (380, 620, 800)
    hint = length_hint(26, printed_start=425, printed_end=450, body_words=9000)
    assert "26 PDF pages" in hint
    assert "380–620" in hint
    assert "30–45" in hint
    assert "35–70" in hint
    assert "60–110" in hint
    assert "printed 425–450" in hint
    assert "three questions" in hint
    assert "length-note" in hint
    assert "Squire" in hint
    assert "see you next time" in hint
    assert "no preamble" in hint
    assert "Do not write a Reflection" in hint
    note = summary_length_note(849, 16)
    assert "over max" in note
    assert "no reason given" in note
    assert "before the closing quote" in note
    assert "13–24 pages" in note
    on_target = summary_length_note(450, 16)
    assert "on target" in on_target
    assert "before the closing quote" in on_target
    assert "long" in summary_length_note(600, 16)
    assert "short" in summary_length_note(328, 16)
    justified = summary_length_note(
        328,
        16,
        justification="literary-structure study of one chapter, like Squire on Alma 63",
    )
    assert "heads-up" in justified
    assert "over max" not in justified
    assert "literary-structure" in justified
    assert "heads-up" in summary_length_note(
        1101,
        24,
        justification="nine parallel points, about 100 words each",
    )
    assert "on target" in qa_length_note(_qa([40, 55, 48]))
    assert "short" in qa_length_note(_qa([40, 20, 48]))
    assert "long" in qa_length_note(_qa([40, 80, 48]))
    assert "over max" in qa_length_note(_qa([40, 95, 48]))
    assert "3 questions" in qa_length_note(_qa([40, 55, 48]))
    assert "expected 3" in qa_length_note(_qa([40, 55, 48], extra_question=True))
    assert "on target" in takeaway_length_note(38)
    assert "long" in takeaway_length_note(50)
    assert "over max" in takeaway_length_note(60)
    assert "short" in takeaway_length_note(22)
    assert "on target" in closing_quote_length_note(90)
    assert "short" in closing_quote_length_note(40)
    assert "over max" in closing_quote_length_note(140)
    assert "on target" in video_length_note(10)
    assert "short" in video_length_note(4)


def test_question_and_video_counts():
    assert question_count(SAMPLE_SUMMARY_MARKDOWN) == 3
    sections = parse_sections(SAMPLE_SUMMARY_MARKDOWN)
    assert video_row_count(sections["video_script"]) == 10


def test_parse_sections_for_current_format():
    sections = parse_sections(SAMPLE_SUMMARY_MARKDOWN)
    assert sections["title"] == "Interpreting Interpreter: Tokens, Not Tonnage"
    assert sections["intro"] == ""
    assert sections["takeaway"].startswith("Scholar argues")
    assert "Fragment W" in sections["qa"]
    assert "warehouse fragment" in sections["summary"]
    assert "see you next time" in sections["video_script"]
    assert "reflection" not in sections
    assert locator_count(SAMPLE_SUMMARY_MARKDOWN) == 7
    assert locator_count('[Lamanite politics] (link to “The present article”; page 426)') == 1
    assert locator_count('**[The verb]** (link to "Three features"; page 2)') == 1


def test_summary_prose_drops_closing_quote_before_the_word_count():
    summary = """
In this article, Derek J. Squire continues to mine Alma.

As Squire concludes (link to "Why does any"; page 83):

> Alma 63 could easily have been written without all the bells and whistles.
"""
    prose = summary_prose(summary)
    assert "mine Alma" in prose
    assert "bells and whistles" not in prose
    assert "As Squire concludes" in prose


def test_extract_length_note_leaves_the_post_and_keeps_the_reason():
    raw = (
        "# Interpreting Interpreter: Alma 63\n\n"
        "## The Takeaway\n\n"
        "Squire outlines a chiasm.\n\n"
        "<!-- length-note: below the band because this is a literary-structure study of one chapter -->\n"
    )
    cleaned, reason = extract_length_note(raw)
    assert "length-note" not in cleaned
    assert cleaned.startswith("# Interpreting Interpreter: Alma 63")
    assert reason is not None
    assert "literary-structure" in reason
    headed = raw + "\n## Length note\n\nAbove the band because of nine parallels.\n"
    cleaned, reason = extract_length_note(headed)
    assert "## Length note" not in cleaned
    assert "nine parallels" in (reason or "")


def test_readable_text_rejoins_line_break_hyphenation():
    assert readable_text("writ -\nten") == "written"
    assert readable_text("writ-\nten") == "written"
    assert readable_text("sophisti -\ncated") == "sophisticated"
    assert readable_text("thirty-seventh") == "thirty-seventh"
    assert readable_text("“written”") == '"written"'
    assert readable_text("hand —the") == "hand-the"
    assert readable_text("hand — the") == "hand-the"
    assert normalize_for_match("Hand — The") == normalize_for_match("hand-the")


def test_strip_preamble_cuts_chatter_glued_to_the_title():
    raw = (
        "I'll pull exact paragraph openings, printed pages, and the closing quotation."
        "# Interpreting Interpreter: Alma 63\n\n"
        "## The Takeaway\n\n"
        "Squire outlines a chiasm in Alma 63.\n"
    )
    cleaned = strip_preamble(raw)
    assert cleaned.startswith("# Interpreting Interpreter: Alma 63")
    assert "I'll pull" not in cleaned
    sections = parse_sections(cleaned)
    assert sections["title"] == "Interpreting Interpreter: Alma 63"
    document = Document(BytesIO(markdown_to_docx(cleaned)))
    assert "I'll pull" not in document.paragraphs[0].text
    assert document.paragraphs[0].text.startswith("Interpreting Interpreter: Alma 63")

    lined = "I'll pull exact paragraph openings.\n\n# Interpreting Interpreter: Alma 63\n"
    assert strip_preamble(lined).startswith("# Interpreting Interpreter: Alma 63")
    numbered = "Look at page # 1 first.# Interpreting Interpreter: Alma 63\n\n## The Takeaway\n\nSquire outlines it.\n"
    assert strip_preamble(numbered).startswith("# Interpreting Interpreter: Alma 63")


def test_strip_reflection_section_drops_retired_block():
    text = """
## The Summary

Keep the summary.

## The Reflection

I like the caution.

## Video Script

Keep the script.
"""
    stripped = strip_reflection_section(text)
    assert "Keep the summary" in stripped
    assert "I like the caution" not in stripped
    assert "The Reflection" not in stripped
    assert "Keep the script" in stripped


def test_markdown_to_docx_matches_template_shape():
    document = Document(BytesIO(markdown_to_docx(SAMPLE_SUMMARY_MARKDOWN)))
    paragraphs = [paragraph.text for paragraph in document.paragraphs]
    assert paragraphs[0].startswith("Interpreting Interpreter: Tokens, Not Tonnage")
    title_runs = document.paragraphs[0].runs
    assert any(run.italic and run.text == "Interpreter" for run in title_runs)
    assert any(run.bold and run.font.size and run.font.size.pt == 16 for run in title_runs)
    assert "The Takeaway" in paragraphs
    assert "The Q&A" in paragraphs
    assert "The Summary" in paragraphs
    assert "Video Script" in paragraphs
    assert "The Reflection" not in paragraphs
    question = next(paragraph for paragraph in document.paragraphs if paragraph.text.startswith("What is Fragment"))
    assert all((run.bold and run.italic) for run in question.runs if run.text.strip())
    video = next(paragraph for paragraph in document.paragraphs if paragraph.text == "Video Script")
    assert all(run.italic and not run.bold for run in video.runs if run.text.strip())
    assert len(document.tables) == 1
    table = document.tables[0]
    assert [cell.text for cell in table.rows[0].cells] == ["#", "Text", "Image"]
    assert len(table.rows) == 11
    assert "see you next time" in table.rows[-1].cells[1].text
    assert table.rows[-1].cells[2].text == "Title page"
    bullet = next(paragraph for paragraph in document.paragraphs if paragraph.text.startswith("[The verb]"))
    assert any(run.bold and run.text == "[The verb]" for run in bullet.runs)
    section = document.sections[0]
    assert section.left_margin.inches == 1
    assert document.core_properties.author == "Kyler Rasmussen"


def _corpus_from_pages(pages: dict[int, str], *, used_printed: bool = True) -> PdfCorpus:
    return PdfCorpus(
        pages=[
            PdfPage(
                viewer_page=index,
                journal_page=journal,
                text=text,
                normalized=normalize_for_match(text),
            )
            for index, (journal, text) in enumerate(pages.items(), start=1)
        ],
        used_printed_pages=used_printed,
    )


def test_verify_locators_keeps_corrects_and_drops():
    corpus = _corpus_from_pages(
        {
            426: "The present article continues an earlier study of Nephite politics.",
            427: "The earliest Book of Mormon references to Lamanite politics are thin.",
            449: "The Book of Mormon does not provide a continuous narrative.",
        }
    )
    markdown = """
Hudson looks at [Lamanite politics] (link to "The present article"; page 426).
He also notes early glimpses (link to "The earliest Book"; page 1).
A fake claim appears (link to "No such paragraph exists here"; page 440).
As he concludes (link to “The Book of Mormon does not”; page 449):
"""
    report = verify_locators(markdown, corpus)
    assert report.kept == 2
    assert report.corrected == 1
    assert report.dropped == 1
    assert '(link to "The present article"; page 426)' in report.markdown
    assert '(link to "The earliest Book"; page 427)' in report.markdown
    assert "page 1" not in report.markdown
    assert "No such paragraph exists here" not in report.markdown
    assert "A fake claim appears." in report.markdown
    assert locator_count(report.markdown) == 3


def test_verify_locators_corrects_a_label_woven_into_the_sentence():
    corpus = _corpus_from_pages(
        {
            70: "Helaman, a son of Alma, takes the records at the close of the war.",
        }
    )
    markdown = (
        'After [briefly summarizing] (link to "Helaman, a son"; page 99) '
        "the chapter, he outlines the chiasm."
    )
    report = verify_locators(markdown, corpus)
    assert report.corrected == 1
    assert report.dropped == 0
    assert 'After [briefly summarizing] (link to "Helaman, a son"; page 70) the chapter' in report.markdown
    assert "page 99" not in report.markdown


def test_verify_locators_keeps_bold_labels_and_corrects_their_pages():
    corpus = _corpus_from_pages(
        {
            2: "Three features of the surrounding clauses matter for the reading.",
        }
    )
    markdown = '**[The verb]** (link to "Three features"; page 9). The verb means confirm.'
    report = verify_locators(markdown, corpus)
    assert report.corrected == 1
    assert '**[The verb]** (link to "Three features"; page 2)' in report.markdown


def test_verify_locators_in_qa_section():
    corpus = _corpus_from_pages({4: "The Bible says nothing about his childhood."})
    markdown = """
## The Q&A

### What do the extra-biblical texts add?

They fill the silence (link to "The Bible says nothing"; page 1).
"""
    report = verify_locators(markdown, corpus)
    assert report.corrected == 1
    assert '(link to "The Bible says nothing"; page 4)' in report.markdown
    assert "What do the extra-biblical texts add?" in report.markdown


def test_verify_locators_does_not_guess_ambiguous_hits():
    corpus = _corpus_from_pages(
        {
            425: "The aftermath of the schism is narrated immediately.",
            440: "The aftermath of the massacre reshaped the coalition.",
            445: "The aftermath of Ammoron's death weakened the war effort.",
        }
    )
    markdown = 'Later (link to "The aftermath of"; page 99) the coalition broke.'
    report = verify_locators(markdown, corpus)
    assert report.dropped == 1
    assert report.corrected == 0
    assert "link to" not in report.markdown


def test_verify_locators_keeps_claimed_page_among_ambiguous_hits():
    corpus = _corpus_from_pages(
        {
            425: "The aftermath of the schism is narrated immediately.",
            440: "The aftermath of the massacre reshaped the coalition.",
        }
    )
    markdown = 'Later (link to "The aftermath of"; page 440) the coalition broke.'
    report = verify_locators(markdown, corpus)
    assert report.kept == 1
    assert '(link to "The aftermath of"; page 440)' in report.markdown


def test_verify_quotes_keeps_exact_and_ellipsis_spans():
    corpus = _corpus_from_pages(
        {
            449: (
                "A close reading of Mosiah through Helaman suggests that Lamanite "
                "history was shaped by more than inherited hostility toward the Nephites. "
                "The text instead reveals a dynamic society whose institutions evolved."
            ),
            450: (
                "The result is a more historically textured picture of the Lamanites, "
                "not merely a static foil to the Nephites."
            ),
        }
    )
    markdown = """
As Hudson concludes:

> A close reading of Mosiah through Helaman suggests that Lamanite history was shaped by more than inherited hostility toward the Nephites. ... The result is a more historically textured picture of the Lamanites, not merely a static foil to the Nephites.
"""
    report = verify_quotes(markdown, corpus)
    assert report.kept == 1
    assert report.corrected == 0
    assert report.dropped == 0
    assert "As Hudson concludes:" in report.markdown


def test_verify_quotes_repairs_paraphrase_using_pdf_wording():
    corpus = _corpus_from_pages(
        {
            3: (
                "As I conclude: although copper can be cargo, in this fragment it more "
                "plausibly memorializes a promise."
            )
        }
    )
    markdown = """
As Scholar concludes:

> although copper can be cargo, in this fragment it obviously memorializes a promise.
"""
    report = verify_quotes(markdown, corpus)
    assert report.corrected == 1
    assert "obviously" not in report.markdown
    assert "more plausibly memorializes a promise" in report.markdown
    assert report.markdown.strip().startswith("> ") or "> although copper can be cargo" in report.markdown


def test_verify_quotes_keeps_a_hyphenated_closing_quote():
    """The Squire PDF splits words across line breaks ('writ -\\nten') and a page header."""
    corpus = _corpus_from_pages(
        {
            83: (
                "Why does any of this matter? Alma 63 could easily have been "
                "writ -\n"
                "ten without all the bells and whistles of literary devices and "
                "sophisti -\n"
                "cated structural techniques. However, the incredible attention "
                "given to"
            ),
            84: (
                "84 • Interpreter 70 (2027)\n"
                "form and style, structure and content, is consistent with what is found "
                "throughout the book of Alma. A middle sentence that the quote will skip. "
                "Mormon appears to give particular care to crafting conclusions. "
                "Even mundane details are shared to match the project at hand —the writing of "
                "scrip -\n"
                "ture."
            ),
        }
    )
    markdown = """
As Squire concludes (link to "Why does any"; page 83):

> Why does any of this matter? Alma 63 could easily have been written without all the bells and whistles of literary devices and sophisticated structural techniques. However, the incredible attention given to form and style, structure and content, is consistent with what is found throughout the book of Alma. .... Mormon appears to give particular care to crafting conclusions. Even mundane details are shared to match the project at hand—the writing of scripture.
"""
    report = verify_quotes(markdown, corpus)
    assert report.kept == 1
    assert report.dropped == 0
    assert "As Squire concludes" in report.markdown
    assert "written without" in report.markdown
    assert "sophisticated" in report.markdown

    bad = """
As Squire concludes:

> Alma 63 could easily have been written by a later editor who invented the chiasm.
"""
    dropped = verify_quotes(bad, corpus)
    assert dropped.dropped == 1
    assert "later editor" not in dropped.markdown
    assert "As Squire concludes" not in dropped.markdown


def test_verify_quotes_drops_unverified_quote_and_attribution():
    corpus = _corpus_from_pages({1: "The warehouse fragment lists ten ingots of bright copper."})
    markdown = """
The fragment lists ingots.

As Scholar concludes:

> This sentence does not appear anywhere in the article at all.

The next paragraph follows.
"""
    report = verify_quotes(markdown, corpus)
    assert report.dropped == 1
    assert "This sentence does not appear" not in report.markdown
    assert "As Scholar concludes:" not in report.markdown
    assert "The fragment lists ingots." in report.markdown
    assert "The next paragraph follows." in report.markdown
