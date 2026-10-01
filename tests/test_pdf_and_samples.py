from pathlib import Path

import pytest

from interpreter_summary.locators import locator_count, verify_locators
from interpreter_summary.quotes import verify_quotes
from interpreter_summary.pdf_utils import (
    PdfError,
    extract_pdf_corpus,
    infer_journal_pages,
    parse_printed_page,
    pdf_page_count,
    validate_pdf,
)
from interpreter_summary.samples import generate_samples, write_sample_pdf
from interpreter_summary.style import load_style_text


def test_validate_pdf_rejects_non_pdf():
    with pytest.raises(PdfError, match="PDF"):
        validate_pdf(b"not a pdf", "notes.txt", 1024)


def test_validate_pdf_rejects_empty():
    with pytest.raises(PdfError, match="empty"):
        validate_pdf(b"", "article.pdf", 1024)


def test_validate_pdf_rejects_oversize():
    with pytest.raises(PdfError, match="larger"):
        validate_pdf(b"%PDF-1.4 oversized", "article.pdf", 8)


def test_sample_pdf_and_style_docx(tmp_path: Path):
    pdf_path, docx_path = generate_samples(tmp_path)
    data = pdf_path.read_bytes()
    validate_pdf(data, pdf_path.name, 5 * 1024 * 1024)
    assert pdf_page_count(data) == 3

    style = load_style_text(docx_path)
    assert "The Takeaway" in style
    assert "The Q&A" in style
    assert "link to" in style
    assert "Video Script" in style
    assert "see you next time" in style
    assert "The Reflection" not in style

    from docx import Document

    document = Document(docx_path)
    assert len(document.tables) == 1
    assert [cell.text for cell in document.tables[0].rows[0].cells] == ["#", "Text", "Image"]


def test_write_sample_pdf_has_printed_page_markers(tmp_path: Path):
    from interpreter_summary.pdf_utils import extract_pdf_text

    path = write_sample_pdf(tmp_path / "article.pdf")
    text = extract_pdf_text(path.read_bytes())
    assert "[Page 1]" in text
    assert "bright copper" in text


def test_parse_printed_page_from_interpreter_headers():
    assert parse_printed_page("426 • Interpreter 69 (2026)\nBody text") == 426
    assert parse_printed_page('Hudson, “Dynastic Dynamics II” • 427\nBody') == 427
    assert parse_printed_page("[Page 3]\nThe warehouse fragment") == 3
    assert parse_printed_page("Dynastic Dynamics II: Lamanite\nAbstract:") is None


def test_infer_journal_pages_fills_unlabeled_first_leaf():
    pages, used = infer_journal_pages([None, 426, 427, 428])
    assert used is True
    assert pages == [425, 426, 427, 428]


def test_sample_pdf_corpus_uses_page_markers(tmp_path: Path):
    path = write_sample_pdf(tmp_path / "article.pdf")
    corpus = extract_pdf_corpus(path.read_bytes())
    assert corpus.used_printed_pages
    assert [page.journal_page for page in corpus.pages] == [1, 2, 3]
    assert corpus.pages_containing("The warehouse fragment") == [1]
    hint = corpus.pagination_hint()
    assert "1–3" in hint


def test_hudson_samples_journal_pages_and_locators():
    root = Path(__file__).resolve().parents[1]
    pdf_path = root / "samples" / "69_16_Hudson.pdf"
    docx_path = root / "samples" / "69_16 (Aug26) Hudson.docx"
    if not pdf_path.exists() or not docx_path.exists():
        pytest.skip("Hudson sample files are not present")

    corpus = extract_pdf_corpus(pdf_path.read_bytes())
    assert corpus.used_printed_pages
    assert corpus.journal_start == 425
    assert corpus.journal_end == 450
    assert corpus.pages[0].viewer_page == 1
    assert "425" not in corpus.pages[0].text
    assert corpus.pages_containing("The earliest Book") == [427]

    style = load_style_text(docx_path)
    report = verify_locators(style, corpus)
    assert locator_count(style) == 19
    assert report.dropped == 0
    assert report.locator_count == 19

    from docx import Document

    paragraphs = [p.text.strip() for p in Document(docx_path).paragraphs if p.text.strip()]
    hudson_quote = next(p for p in paragraphs if p.startswith("A close reading of Mosiah"))
    quote_report = verify_quotes(
        f"As Hudson concludes:\n\n> {hudson_quote}\n",
        corpus,
    )
    assert quote_report.dropped == 0
    assert quote_report.quote_count == 1
