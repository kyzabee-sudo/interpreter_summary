"""Regression checks from Kyler's proofread of the Smith draft, plus the pipeline."""

import json
from io import BytesIO

import httpx
import pytest
from docx import Document
from fpdf import FPDF

from interpreter_summary.checks import (
    inspect_draft,
    model_finding_resolved,
    quotation_findings,
    remaining_locator_findings,
    render_check_report,
    surface_findings,
    tidy_draft,
    video_findings,
)
from interpreter_summary.cli import main
from interpreter_summary.config import Settings
from interpreter_summary.export import markdown_to_docx, parse_sections
from interpreter_summary.grok import GrokClient
from interpreter_summary.length import count_words, summary_prose, summary_word_target
from interpreter_summary.locators import LocatorReport, verify_locators
from interpreter_summary.pdf_utils import (
    PdfCorpus,
    PdfPage,
    extract_pdf_corpus,
    normalize_for_match,
)
from interpreter_summary.quotes import QuoteReport
from interpreter_summary.review import (
    REVIEW_SYSTEM_PROMPT,
    build_review_user_prompt,
    parse_model_findings,
)
from interpreter_summary.samples import SAMPLE_SUMMARY_MARKDOWN, write_sample_pdf
from interpreter_summary.service import SummaryResult, summarize_pdf


SMITH_PAGE_155 = (
    "If Nephi's prologue establishes the pattern of deliverance, the same framing "
    "may extend through the rest of 1 Nephi and possibly beyond. "
    "Reynolds's model stops at the opening vision. "
    "The five-part chiasm mislocates the center of Nephi's account. "
    "Robb Smith finds the pattern in 1 Nephi 1:1–3 only."
)
SMITH_PAGE_156 = (
    "The lands at E3 sit inside the judgment frame of Lehi's vision. "
    "A later note treats the opening frame as local to that vision."
)


def _corpus(pages: dict[int, str]) -> PdfCorpus:
    built = []
    for index, (journal, text) in enumerate(pages.items(), start=1):
        built.append(
            PdfPage(
                viewer_page=index,
                journal_page=journal,
                text=text,
                normalized=normalize_for_match(text),
            )
        )
    return PdfCorpus(pages=built, used_printed_pages=True)


def _smith_corpus() -> PdfCorpus:
    return _corpus({155: SMITH_PAGE_155, 156: SMITH_PAGE_156})


def _codes(findings) -> set[str]:
    return {finding.code for finding in findings}


def _pdf(pages: list[tuple[int, str]]) -> bytes:
    pdf = FPDF()
    pdf.set_auto_page_break(auto=True, margin=15)
    for number, text in pages:
        pdf.add_page()
        pdf.set_font("Times", size=12)
        safe = (
            text.replace("\u2013", "-")
            .replace("\u2014", "-")
            .replace("\u2018", "'")
            .replace("\u2019", "'")
            .replace("\u201c", '"')
            .replace("\u201d", '"')
        )
        pdf.multi_cell(0, 6, f"[Page {number}]\n\n{safe}")
    return bytes(pdf.output())


def test_locator_phrase_must_be_on_the_stated_journal_page():
    corpus = _smith_corpus()
    # Line-break hyphenation and a curly apostrophe still count as the phrase.
    hyphenated = _corpus(
        {
            155: "If Nephi’s pro-\nlogue establishes the pattern of deliverance, the frame holds.",
        }
    )
    assert hyphenated.phrase_pages('If Nephi\'s prologue establishes') == [155]

    draft = (
        'After [lands at E3] (link to "If Nephi\'s prologue establishes"; page 156) '
        "the vision, the frame is set."
    )
    raw = remaining_locator_findings(draft, corpus)
    assert len(raw) == 1
    assert raw[0].page == 155
    assert "156" in raw[0].problem
    assert "155" in raw[0].problem

    report = verify_locators(draft, corpus)
    assert report.corrected == 1
    assert '(link to "If Nephi\'s prologue establishes"; page 155)' in report.markdown
    audit = inspect_draft(
        report.markdown,
        corpus,
        article_pages=2,
        locator_report=report,
    )
    assert any("155" in note and "156" in note for note in audit.fixed_notes)
    assert "locator-page" not in _codes(audit.open_findings)
    assert corpus.journal_offset() == 154


def test_pdf_offset_is_derived_from_printed_pages():
    data = _pdf([(155, SMITH_PAGE_155), (156, SMITH_PAGE_156)])
    corpus = extract_pdf_corpus(data)
    assert corpus.used_printed_pages is True
    assert corpus.journal_start == 155
    assert corpus.journal_offset() == 154
    assert corpus.phrase_pages("If Nephi's prologue establishes") == [155]
    assert "lands at E3" in corpus.pages[1].text or corpus.phrase_pages("The lands at E3") == [156]


def test_altered_quote_is_flagged_with_the_article_wording():
    corpus = _smith_corpus()
    draft = (
        "## The Summary\n\n"
        'Reynolds’s model “mislocate the center of Nephi’s account” and then stops.\n'
    )
    findings = quotation_findings(draft, corpus)
    altered = [finding for finding in findings if finding.code == "quote-altered"]
    assert altered
    found = altered[0]
    assert found.page == 155
    assert "mislocates" in found.evidence
    assert "five-part chiasm" in found.evidence
    assert "mislocate" in found.problem


def test_surface_junk_and_tidy_fixes_ranges_and_apostrophes():
    draft = (
        "## The Takeaway\n\n"
        "Smith sees the Nephi's final declaration  and then stops\n\n"
        "## The Summary\n\n"
        "In this article, from . Robb Smith  argues the point (1 Nephi 1:1-3 and 4–7).\n\n"
        "The the appendices weigh the evidence XXXXX [] and a (stray parenthesis.\n\n"
        'He said "hello.\n'
    )
    raw = surface_findings(draft)
    codes = _codes(raw)
    assert "placeholder" in codes
    assert "doubled-word" in codes
    assert "dropped-initial" in codes
    assert "article-possessive" in codes
    assert "dash-range" in codes
    assert "apostrophe" in codes
    assert "unbalanced" in codes
    assert "terminal-punct" in codes
    assert any("mixed" in finding.problem.lower() or "hyphen" in finding.problem.lower() for finding in raw)
    assert not any("Lord's" in finding.problem for finding in surface_findings("in the Lord's house."))

    cleaned, notes = tidy_draft(draft)
    assert "1–3" in cleaned
    assert "1-3" not in cleaned
    assert "4–7" in cleaned
    assert "Nephi’s" in cleaned
    assert "  " not in cleaned
    assert any("en dash" in note for note in notes)
    assert any("apostrophe" in note for note in notes)
    assert "XXXXX" in cleaned
    assert "dash-range" not in _codes(surface_findings(cleaned))
    assert "apostrophe" not in _codes(surface_findings(cleaned))


def test_summary_count_excludes_locator_parentheticals_and_keeps_the_band():
    base = "In this article, Robb Smith argues the frame holds and then he stops."
    with_locator = (
        'In this article, Robb Smith argues the frame holds '
        '(link to "If Nephi\'s prologue establishes the pattern"; page 155) '
        "and then he stops."
    )
    assert count_words(base) == count_words(with_locator)
    assert count_words(with_locator) < count_words(
        with_locator.replace("(link to", "link to")
    )
    # The existing page-scaled band is unchanged. A short article is not forced to 250–400.
    assert summary_word_target(8) == (200, 330, 400)
    summary = (
        "## The Summary\n\n"
        f"{with_locator}\n\n"
        'As Smith concludes (link to "may extend"; page 155):\n\n'
        "> may extend through the rest of 1 Nephi and possibly beyond.\n"
    )
    prose = summary_prose(parse_sections(summary)["summary"])
    assert count_words(prose) == count_words(base + " As Smith concludes:")
    audit = inspect_draft(summary, _smith_corpus(), article_pages=8)
    assert any("not counting locator parentheticals" in line for line in audit.count_lines)
    assert any("200–330" in finding.problem or "200-330" in finding.problem for finding in audit.open_findings)


def test_takeaway_band_and_script_completeness():
    short = "## The Takeaway\n\nSmith argues the point.\n"
    audit = inspect_draft(short, _smith_corpus(), article_pages=3)
    assert any(finding.code == "takeaway-length" for finding in audit.open_findings)
    assert any("20–35" in finding.problem or "20-35" in finding.problem for finding in audit.open_findings)

    rows = "\n".join(
        f"| {index} | Point {index} is new. | Cue {index} |" for index in range(1, 11)
    )
    good = (
        "## Video Script\n\n"
        "| # | Text | Image |\n| --- | --- | --- |\n"
        + rows.replace(
            "| 10 | Point 10 is new. | Cue 10 |",
            "| 10 | Check the article, and I'll see you next time. | Title page |",
        )
    )
    assert video_findings(good) == []
    curly = good.replace("I'll", "I’ll")
    assert video_findings(curly) == []

    short_script = good.replace("| 9 | Point 9 is new. | Cue 9 |\n", "").replace(
        "| 8 | Point 8 is new. | Cue 8 |\n", ""
    )
    assert any(finding.code == "script-rows" for finding in video_findings(short_script))
    missing_signoff = good.replace("and I'll see you next time.", "thanks for reading.")
    assert any(finding.code == "script-signoff" for finding in video_findings(missing_signoff))
    empty_cell = good.replace("| Cue 4 |", "|  |")
    assert any(finding.code == "script-cells" for finding in video_findings(empty_cell))
    assert any(finding.code == "script-table" for finding in video_findings("## Video Script\n\nNo table."))


def test_model_findings_round_trip_the_smith_contradictions():
    lowered = REVIEW_SYSTEM_PROMPT.lower()
    assert "some scholars" in lowered
    assert "mislocates" in lowered
    assert "may extend" in lowered
    assert "from . robb smith" in lowered
    assert "the nephi" in lowered
    assert "appendices weigh" in lowered
    assert "hedg" in lowered
    prompt = build_review_user_prompt(
        pagination="Printed journal pages are 155–156.",
        article_text="[Page 155]\n" + SMITH_PAGE_155,
        article_truncated=False,
        draft="draft",
        open_findings=[],
    )
    assert "155" in prompt
    assert "If Nephi's prologue" in prompt

    raw = """```json
    {"findings": [
      {
        "location": "The Q&A, question 2",
        "problem": "The answer says the pattern \\"doesn't appear to be repeated.\\"",
        "evidence": "may extend through the rest of 1 Nephi and possibly beyond",
        "page": 155,
        "suggested_fix": "Keep the author's may and possibly."
      },
      {
        "location": "The Q&A",
        "problem": "The draft says Smith claims both models mislocate the center.",
        "evidence": "The five-part chiasm mislocates the center of Nephi's account.",
        "page": "155",
        "suggested_fix": "Attribute mislocates to the five-part chiasm only."
      }
    ]}
    ```"""
    findings, error = parse_model_findings(raw)
    assert error is None
    assert len(findings) == 2
    assert findings[0].page == 155
    assert findings[0].source == "model"
    assert "possibly beyond" in findings[1].evidence or "mislocates" in findings[1].evidence
    report = render_check_report(
        corpus=_smith_corpus(),
        review_status="ran",
        revision_status="applied once",
        fixed_notes=[],
        open_findings=findings,
        count_lines_text=["Takeaway: 28 words."],
    )
    assert "journal page 155" in report
    assert "may extend through the rest of 1 Nephi and possibly beyond" in report
    assert "five-part chiasm mislocates" in report
    assert "Suggested fix" in report
    assert "PDF viewer page + 154" in report
    assert model_finding_resolved(findings[0], "The pattern may extend and possibly beyond.")
    assert not model_finding_resolved(findings[0], 'The pattern "doesn\'t appear to be repeated."')
    assert parse_model_findings('{"findings": []}') == ([], None)


def _handler(routes: dict[str, str]):
    def respond(request: httpx.Request) -> httpx.Response:
        if request.method == "POST" and request.url.path == "/v1/files":
            return httpx.Response(200, json={"id": "file-123"})
        if request.method == "DELETE":
            return httpx.Response(200, json={"deleted": True})
        if request.method == "POST" and request.url.path == "/v1/responses":
            payload = json.loads(request.content.decode() or "{}")
            routes.setdefault("seen", []).append(payload)
            system = ""
            for item in payload.get("input") or []:
                if item.get("role") == "system":
                    system = item.get("content") or ""
            if "Return only JSON" in system:
                text = routes["review"]
            elif "revised post" in system:
                text = routes["revision"]
            else:
                text = routes["draft"]
            event = json.dumps({"type": "response.output_text.delta", "delta": text})
            return httpx.Response(200, text=f"data: {event}\n\ndata: [DONE]\n\n")
        return httpx.Response(404, json={"error": "unhandled"})

    return respond


def _system_text(payload: dict) -> str:
    for item in payload.get("input") or []:
        if item.get("role") == "system":
            content = item.get("content") or ""
            return content if isinstance(content, str) else ""
    return ""


def _file_attached(payload: dict) -> bool:
    for item in payload.get("input") or []:
        content = item.get("content")
        if not isinstance(content, list):
            continue
        for part in content:
            if isinstance(part, dict) and part.get("type") == "input_file":
                return True
    return False


@pytest.mark.asyncio
async def test_pipeline_revises_once_without_reattaching_the_pdf(tmp_path):
    pdf_path = write_sample_pdf(tmp_path / "article.pdf")
    bad = SAMPLE_SUMMARY_MARKDOWN.replace(
        "Scholar treats the analogy as suggestive rather than decisive",
        'Scholar says the pattern "doesn\'t appear to be repeated"',
    ).replace(
        "In this article, A. Sample Scholar",
        'In this article, A. Sample Scholar (link to "No such paragraph exists here"; page 9)',
    )
    review = json.dumps(
        {
            "findings": [
                {
                    "location": "The Q&A",
                    "problem": 'The answer says the pattern "doesn\'t appear to be repeated."',
                    "evidence": "may extend through the rest of 1 Nephi and possibly beyond",
                    "page": 155,
                    "suggested_fix": "Keep may and possibly.",
                }
            ]
        }
    )
    routes = {"draft": bad, "review": review, "revision": SAMPLE_SUMMARY_MARKDOWN, "seen": []}
    settings = Settings(xai_api_key="test-key", xai_model="grok-4.6")
    transport = httpx.MockTransport(_handler(routes))
    async with httpx.AsyncClient(transport=transport, base_url="https://api.x.ai") as raw:
        client = GrokClient(settings, client=raw)
        result = await summarize_pdf(
            pdf_path.read_bytes(),
            "article.pdf",
            settings,
            client=client,
        )
    seen = routes["seen"]
    assert len(seen) == 4
    assert _file_attached(seen[0])
    assert "reasoning" not in seen[0]
    assert not _file_attached(seen[1])
    assert not _file_attached(seen[2])
    assert not _file_attached(seen[3])
    assert seen[1]["reasoning"] == {"effort": "low"}
    assert seen[2]["reasoning"] == {"effort": "low"}
    assert seen[3]["reasoning"] == {"effort": "low"}
    assert "Output only the Lay Academic script" in _system_text(seen[3])
    assert "Lay Academic" not in result.markdown
    assert "doesn't appear to be repeated" not in result.markdown
    assert "Check report" not in result.markdown
    assert "may extend through the rest of 1 Nephi and possibly beyond" in result.check_report
    assert "applied once" in result.check_report
    assert "No such paragraph exists here" in result.check_report
    assert "not searched again" in result.check_report
    document = Document(BytesIO(markdown_to_docx(result.markdown)))
    blob = "\n".join(paragraph.text for paragraph in document.paragraphs)
    assert "Check report" not in blob
    assert "Still needs an eye" not in blob
    for table in document.tables:
        for row in table.rows:
            for cell in row.cells:
                assert "Check report" not in cell.text


@pytest.mark.asyncio
async def test_no_verify_is_offline_and_still_flags_the_smith_locator(tmp_path):
    data = _pdf([(155, SMITH_PAGE_155), (156, SMITH_PAGE_156)])
    draft = """# Interpreting Interpreter: Framing the Vision

## The Takeaway

Smith argues that Nephi’s prologue frames judgment and deliverance in the opening of 1 Nephi.

## The Q&A

### What does the prologue do?

It sets a frame of judgment and deliverance for the opening vision, and Smith keeps the author’s tentative tone.

### Where is the center?

The five-part chiasm is the structure that misplaces its own middle, while Reynolds’s model stops earlier.

### What should a reader watch?

The opening frame is doing structural work, and the article stays inside 1 Nephi 1 rather than the whole book.

## The Summary

In this article, from . Robb Smith reads the opening. After [lands at E3] (link to "If Nephi's prologue establishes"; page 156) the vision, he treats the frame as deliverance. XXXXX

As Smith concludes (link to "may extend through the rest"; page 155):

> may extend through the rest of 1 Nephi and possibly beyond.

## Video Script

| # | Text | Image |
| --- | --- | --- |
| 1 | Nephi opens with a frame. | Prologue |
| 2 | Robb Smith is the one making this case, not a crowd of unnamed scholars. | Title page |
| 3 | The five-part chiasm is the piece that misplaces the center. | Chiasm |
| 4 | Reynolds’s model stops at the opening vision. | Reynolds |
| 5 | The lands at E3 sit inside Lehi’s vision. | Map |
| 6 | The author is tentative about how far the pattern goes. | Hedge |
| 7 | Journal pages are not the PDF viewer index. | Header |
| 8 | One appendix is mentioned and then left alone. | Appendix |
| 9 | The closing line is the author’s, not a slogan. | Quote |
| 10 | Check the article, and I'll see you next time. | Title page |
"""
    seen: list[dict] = []

    def respond(request: httpx.Request) -> httpx.Response:
        if request.method == "POST" and request.url.path == "/v1/files":
            return httpx.Response(200, json={"id": "file-1"})
        if request.method == "DELETE":
            return httpx.Response(200, json={"deleted": True})
        if request.url.path == "/v1/responses":
            seen.append(json.loads(request.content.decode() or "{}"))
            event = json.dumps({"type": "response.output_text.delta", "delta": draft})
            return httpx.Response(200, text=f"data: {event}\n\ndata: [DONE]\n\n")
        return httpx.Response(404)

    settings = Settings(xai_api_key="test-key", xai_model="grok-4.6")
    transport = httpx.MockTransport(respond)
    async with httpx.AsyncClient(transport=transport, base_url="https://api.x.ai") as raw:
        client = GrokClient(settings, client=raw)
        result = await summarize_pdf(
            data,
            "smith.pdf",
            settings,
            client=client,
            model_verify=False,
        )
    assert len(seen) == 2
    assert _file_attached(seen[0])
    assert not _file_attached(seen[1])
    assert seen[1]["reasoning"] == {"effort": "low"}
    assert "Output only the Lay Academic script" in _system_text(seen[1])
    assert "Lay Academic" not in result.markdown
    assert "page 156" not in result.markdown
    assert 'page 155' in result.markdown
    assert "skipped (--no-verify)" in result.check_report
    assert "155" in result.check_report and "156" in result.check_report
    assert "XXXXX" in result.check_report
    assert "from . Robb" in result.markdown
    assert "dropped initial" in result.check_report
    assert "PDF viewer page + 154" in result.check_report
    assert "Check report" not in result.markdown


def test_cli_writes_the_check_report_beside_the_draft_and_not_in_the_docx(tmp_path, monkeypatch, capsys):
    pdf_path = write_sample_pdf(tmp_path / "article.pdf")
    captured: dict = {}

    async def fake_summarize(*args, **kwargs):
        captured["model_verify"] = kwargs["model_verify"]
        return SummaryResult(
            markdown=SAMPLE_SUMMARY_MARKDOWN,
            title="Interpreting Interpreter: Tokens, Not Tonnage",
            intro="",
            takeaway="Scholar argues the point in enough words for a takeaway here.",
            qa="### What?\n\nAn answer.",
            summary="In this article, A. Sample Scholar writes.",
            video_script="| # | Text | Image |",
            page_count=3,
            locator_count=0,
            locators_corrected=0,
            locators_dropped=0,
            quotes_kept=0,
            quotes_corrected=0,
            quotes_dropped=0,
            journal_page_start=1,
            journal_page_end=3,
            used_printed_pages=True,
            model="grok-4.6",
            locator_report=LocatorReport(markdown=""),
            quote_report=QuoteReport(markdown=""),
            summary_word_count=10,
            summary_length_note="summary",
            closing_quote_word_count=0,
            closing_quote_length_note="quote",
            takeaway_word_count=10,
            takeaway_length_note="takeaway",
            qa_word_count=10,
            qa_length_note="qa",
            video_row_count=10,
            video_length_note="video",
            check_report="# Check report\n\nNothing unresolved.\n",
        )

    monkeypatch.setattr("interpreter_summary.cli.summarize_pdf", fake_summarize)
    output = tmp_path / "summary.md"
    docx_path = tmp_path / "summary.docx"
    main(
        [
            "summarize",
            str(pdf_path),
            "-o",
            str(output),
            "--docx",
            str(docx_path),
            "--no-verify",
        ]
    )
    assert captured["model_verify"] is False
    report = tmp_path / "summary.check.md"
    assert report.read_text(encoding="utf-8").startswith("# Check report")
    assert "Check report" not in output.read_text(encoding="utf-8")
    document = Document(docx_path)
    blob = "\n".join(paragraph.text for paragraph in document.paragraphs)
    assert "Check report" not in blob
    assert "Nothing unresolved" not in blob
    assert not (tmp_path / "summary - Lay Academic Script.docx").exists()
    stderr = capsys.readouterr().err
    assert "# Check report" in stderr
