import json
import zipfile
from io import BytesIO
from pathlib import Path
from xml.etree import ElementTree as ET

import httpx
import pytest
from docx import Document
from docx.oxml.ns import qn
from fastapi.testclient import TestClient

from interpreter_summary.cli import main
from interpreter_summary.config import Settings
from interpreter_summary.export import markdown_to_docx
from interpreter_summary.grok import GrokClient
from interpreter_summary.lay_academic import (
    extract_lay_script,
    lay_academic_docx_path,
    lay_academic_to_docx,
    lay_script_notes,
    split_lay_academic,
    strip_lay_author_credit,
)
from interpreter_summary.length import lay_academic_length_hint, length_hint
from interpreter_summary.prompts import (
    LAY_ACADEMIC_SYSTEM_PROMPT,
    SYSTEM_PROMPT,
    build_lay_academic_user_prompt,
    build_user_prompt,
)
from interpreter_summary.samples import SAMPLE_SUMMARY_MARKDOWN, write_sample_pdf
from interpreter_summary.service import summarize_pdf
from interpreter_summary.style import load_style_text
from interpreter_summary.web.app import app

W_NS = "{http://schemas.openxmlformats.org/wordprocessingml/2006/main}"
A_NS = "{http://schemas.openxmlformats.org/drawingml/2006/main}"

LAY_SAMPLE = """\
# Copper Tokens: Lay Academic Script

Alternate script for A. Sample Scholar, “Copper, Covenants, and the Case of the Missing Ingots,” Interpreter 99.

## Video Script (Lay Academic Version)

| # | Narration (spoken) | On-screen caption | Image |
| --- | --- | --- | --- |
| 1 | The historian A. Sample Scholar argues that ten ingots of bright copper in Fragment W are a covenant token, not warehouse cargo. | The historian A. Sample Scholar argues / that ten ingots // of bright copper / in Fragment W / are a covenant token, / not warehouse cargo. | Bright copper ingots on Wikimedia (`01 - Bright Copper.jpg`). |
| 2 | Scholar [reads the list] (link to "ten ingots of"; page 1) as a token rather than as cargo. | Scholar reads the list / as a token / rather than as cargo. | Joseph Smith Papers journal page (`02 - Journal Page.jpg`). |
| 3 | The verb that looks like weigh is the verb the corpus uses when someone confirms a promise. | The verb that looks like weigh / is the verb the corpus uses / when someone confirms a promise. | Article quote card of the weigh passage (`03 - Weigh Quote Card.png`). |
| 4 | The adjective bright keeps company with holiness, not with words for ore. | The adjective bright / keeps company with holiness, / not with words for ore. | Church History Library ore sample (`04 - Ore Sample.jpg`). |
| 5 | The list ends with a witness formula that belongs in a covenant text. | The list ends / with a witness formula / that belongs / in a covenant text. | Witness formula screenshot (`05 - Witness Formula.png`). |
| 6 | A comparative table of treaty deposits pairs metal tokens with oath witnesses. | A comparative table / of treaty deposits / pairs metal tokens / with oath witnesses. | Treaty deposits table from the article (`06 - Treaty Deposits.png`). |
| 7 | Scholar treats the analogy as suggestive, because Fragment W is short. | Scholar treats the analogy / as suggestive, / because Fragment W is short. | Fragment W photograph (`07 - Fragment W.jpg`). |
| 8 | The inventory reading has to explain that witness line as decoration. | The inventory reading / has to explain / that witness line / as decoration. | Shipping receipt comparison (`08 - Shipping Receipt.jpg`). |
| 9 | The closing sentence is Scholar’s, and it keeps the author’s hedge. | The closing sentence / is Scholar’s, / and it keeps / the author’s hedge. | Closing sentence quote card (`09 - Closing Quote Card.png`). |
| 10 | Read Scholar's full article at interpreter foundation dot org. | Read Scholar's full article / at interpreterfoundation.org. | Article title page (`10 - Article Title Page.png`). |

## Post Caption

Interpreting Interpreter: Copper Tokens

## Image Sources

- row 1: 01 - Bright Copper.jpg. https://commons.wikimedia.org/wiki/File:Example.jpg. Credit: Wikimedia Commons (public domain).
- row 2: 02 - Journal Page.jpg. https://www.josephsmithpapers.org/example. Credit: The Joseph Smith Papers / Church History Library.
- row 10: 10 - Article Title Page.png. Article screenshot of the title and abstract. Credit: Scholar, Interpreter 99 (article screenshot).
"""


def test_prompts_and_style_guide_describe_the_lay_script():
    guide = load_style_text()
    for text in (SYSTEM_PROMPT, guide, LAY_ACADEMIC_SYSTEM_PROMPT, lay_academic_length_hint()):
        lowered = text.lower()
        assert "lay academic" in lowered
        assert "narration (spoken)" in lowered
        assert "7" in text
        assert "interpreterfoundation.org" in lowered
        assert "page number" in lowered
        assert "sign-off" in lowered or "sign off" in lowered
    for text in (SYSTEM_PROMPT, guide, LAY_ACADEMIC_SYSTEM_PROMPT):
        lowered = text.lower()
        assert "joseph smith papers" in lowered
        assert "wikimedia" in lowered
        assert "church history library" in lowered
    assert "Do not append a Lay Academic script" in SYSTEM_PROMPT
    assert "separate document" in SYSTEM_PROMPT
    assert "Output only the Lay Academic script" in LAY_ACADEMIC_SYSTEM_PROMPT
    assert "hy-er-AT-ik" in LAY_ACADEMIC_SYSTEM_PROMPT
    assert "eighteen thirty-five" in LAY_ACADEMIC_SYSTEM_PROMPT
    assert " // " in LAY_ACADEMIC_SYSTEM_PROMPT
    assert "credit" in guide.lower()
    assert "NN - Short Desc" in LAY_ACADEMIC_SYSTEM_PROMPT
    assert "You are Kyler Rasmussen" in SYSTEM_PROMPT
    assert "Interpreting Interpreter: [Title] By Kyler Rasmussen" not in SYSTEM_PROMPT
    assert "Interpreting Interpreter: <Topic> By Kyler Rasmussen" not in LAY_ACADEMIC_SYSTEM_PROMPT
    assert "Interpreting Interpreter: Copper Tokens By Kyler Rasmussen" not in LAY_ACADEMIC_SYSTEM_PROMPT
    assert "Interpreting Interpreter: <Topic>" in LAY_ACADEMIC_SYSTEM_PROMPT
    assert "no author" in LAY_ACADEMIC_SYSTEM_PROMPT.lower()
    assert "author line" in LAY_ACADEMIC_SYSTEM_PROMPT.lower()
    lay_guide = guide.split("## Lay Academic script", 1)[1].split("## In-text references", 1)[0]
    assert "`Interpreting Interpreter: <Topic>`" in lay_guide
    assert "By Kyler Rasmussen`" not in lay_guide
    assert "no author" in lay_guide.lower()
    assert "Write as **Kyler Rasmussen**" in guide
    assert "no author" in lay_academic_length_hint().lower()
    lay_user = build_lay_academic_user_prompt(
        pagination="page 1",
        article_text="Article text.",
        article_truncated=False,
        draft="Draft.",
        length_hint_text="exactly 10 rows",
    )
    assert "no author" in lay_user.lower()
    assert "By Kyler Rasmussen" not in lay_user
    readme = Path(__file__).resolve().parents[1].joinpath("README.md").read_text(encoding="utf-8")
    lay_readme = readme.split("## Lay Academic script", 1)[1].split("### Video Images", 1)[0]
    assert "Interpreting Interpreter: [Title]" in lay_readme
    assert "By Kyler Rasmussen" not in lay_readme
    assert "no author" in lay_readme.lower()
    user = build_user_prompt("House style.")
    assert "Do not append a Lay Academic script" in user
    assert "Do not include a video script" not in user
    hint = length_hint(8)
    assert "Lay Academic" in hint
    assert "see you next time" in hint
    assert "interpreterfoundation.org" in hint
    assert "exactly 10 rows" in hint


def test_lay_filename_uses_the_article_stem():
    path = lay_academic_docx_path("articles/70_09 (Nov26) Gee.AI.docx")
    assert path.as_posix() == "articles/70_09 (Nov26) Gee - Lay Academic Script.docx"
    again = lay_academic_docx_path(path)
    assert again.name == path.name
    assert lay_academic_docx_path("summary.md").name == "summary - Lay Academic Script.docx"


def test_split_removes_the_lay_script_from_the_main_post():
    combined = SAMPLE_SUMMARY_MARKDOWN.rstrip() + "\n\n" + LAY_SAMPLE
    main, lay = split_lay_academic(combined)
    assert main.startswith("# Interpreting Interpreter: Tokens, Not Tonnage")
    assert "Lay Academic" not in main
    assert "Post Caption" not in main
    assert "Narration (spoken)" not in main
    assert lay.startswith("# Copper Tokens: Lay Academic Script")
    assert extract_lay_script("Here is the script.\n\n" + LAY_SAMPLE) == LAY_SAMPLE
    assert extract_lay_script(SAMPLE_SUMMARY_MARKDOWN) == ""


def test_caption_notes_flag_pages_small_words_and_split_names():
    assert lay_script_notes(LAY_SAMPLE) == []
    paged = LAY_SAMPLE.replace(
        "not warehouse cargo.",
        "not warehouse cargo on page 12.",
        1,
    )
    assert "page number in the script" in lay_script_notes(paged)
    small = LAY_SAMPLE.replace(
        "The historian A. Sample Scholar argues / that ten ingots // of bright copper",
        "The historian argues for / the",
    )
    notes = " ".join(lay_script_notes(small))
    assert "ends on" in notes
    named = LAY_SAMPLE.replace(
        "The historian A. Sample Scholar argues / that ten ingots // of bright copper",
        "William / W. Phelps wrote it",
    )
    assert any("split a name" in note for note in lay_script_notes(named))


def test_lay_docx_matches_the_four_column_layout():
    data = lay_academic_to_docx(LAY_SAMPLE)
    document = Document(BytesIO(data))
    title = document.paragraphs[0]
    assert title.text.lstrip().startswith("Copper Tokens: Lay Academic Script")
    assert "Interpreting Interpreter" not in title.text
    assert all(not run.italic for run in title.runs)
    assert any(run.bold and run.font.size and run.font.size.pt == 16 for run in title.runs)
    headings = {paragraph.text: paragraph for paragraph in document.paragraphs}
    video = headings["Video Script (Lay Academic Version)"]
    assert all(run.italic and not run.bold for run in video.runs if run.text.strip())
    post = headings["Post Caption"]
    assert all(run.bold and not run.italic for run in post.runs if run.text.strip())
    assert "Image Sources" in headings
    caption = next(
        paragraph.text
        for paragraph in document.paragraphs
        if paragraph.text.startswith("Interpreting Interpreter: Copper Tokens")
    )
    assert caption == "Interpreting Interpreter: Copper Tokens"
    assert "Kyler Rasmussen" not in caption
    assert "By Kyler" not in "\n".join(paragraph.text for paragraph in document.paragraphs)
    assert document.core_properties.author != "Kyler Rasmussen"
    assert b"Kyler Rasmussen" not in data
    assert len(document.tables) == 1
    table = document.tables[0]
    assert [cell.text for cell in table.rows[0].cells] == [
        "#",
        "Narration (spoken)",
        "On-screen caption",
        "Image",
    ]
    assert len(table.rows) == 11
    assert "see you next time" not in table.rows[-1].cells[1].text
    assert "interpreter foundation dot org" in table.rows[-1].cells[1].text
    assert "interpreterfoundation.org" in table.rows[-1].cells[2].text
    assert "page" not in table.rows[-1].cells[2].text.lower()
    spoken = table.rows[2].cells[1].paragraphs[0]
    locator = next(run for run in spoken.runs if "reads the list" in run.text)
    assert not locator.bold
    assert str(locator.font.color.rgb) == "FF0000"
    assert all(not (run.bold and "[" in (run.text or "")) for row in table.rows for cell in row.cells for paragraph in cell.paragraphs for run in paragraph.runs)
    caption_xml = table.rows[1].cells[2]._tc.xml
    assert caption_xml.count("<w:br") + caption_xml.count(":br") >= 1
    bullet = next(paragraph for paragraph in document.paragraphs if "Bright Copper" in paragraph.text)
    assert bullet.style.name == "List Paragraph"
    widths = table._tbl.tblPr.find(qn("w:tblW"))
    assert widths.get(qn("w:w")) == "9001"
    _assert_theme_and_no_page_field(data)


def test_main_ai_docx_stays_summary_only():
    combined = SAMPLE_SUMMARY_MARKDOWN.rstrip() + "\n\n" + LAY_SAMPLE
    main, _lay = split_lay_academic(combined)
    main_doc = Document(BytesIO(markdown_to_docx(main)))
    sample_doc = Document(BytesIO(markdown_to_docx(SAMPLE_SUMMARY_MARKDOWN)))
    assert [paragraph.text for paragraph in main_doc.paragraphs] == [
        paragraph.text for paragraph in sample_doc.paragraphs
    ]
    assert len(main_doc.tables) == 1
    assert [cell.text for cell in main_doc.tables[0].rows[0].cells] == ["#", "Text", "Image"]
    blob = "\n".join(paragraph.text for paragraph in main_doc.paragraphs)
    for cell_text in (cell.text for row in main_doc.tables[0].rows for cell in row.cells):
        blob += "\n" + cell_text
    assert "Lay Academic" not in blob
    assert "Narration (spoken)" not in blob
    assert "Post Caption" not in blob
    assert "Image Sources" not in blob
    assert "see you next time" in main_doc.tables[0].rows[-1].cells[1].text
    assert main_doc.paragraphs[0].text.lstrip().startswith("Interpreting Interpreter: Tokens, Not Tonnage")
    assert main_doc.core_properties.author == "Kyler Rasmussen"
    direct = Document(BytesIO(markdown_to_docx(SAMPLE_SUMMARY_MARKDOWN)))
    assert [cell.text for cell in direct.tables[0].rows[0].cells] == ["#", "Text", "Image"]
    assert len(direct.tables[0].rows) == 11
    assert direct.core_properties.author == "Kyler Rasmussen"
    assert direct.paragraphs[0].text.lstrip().startswith("Interpreting Interpreter: Tokens, Not Tonnage")


def test_lay_output_omits_an_author_credit():
    credited = LAY_SAMPLE.replace(
        "Interpreting Interpreter: Copper Tokens",
        "Interpreting Interpreter: Copper Tokens By Kyler Rasmussen",
    )
    credited = credited.replace(
        "## Image Sources\n",
        "## End card\n\nInterpreting Interpreter: Copper Tokens\n\nBy Kyler Rasmussen\n\n"
        "Author: Kyler Rasmussen\n\n"
        "## Image Sources\n"
        "- Credit: Kyler Rasmussen\n",
    )
    cleaned = strip_lay_author_credit(credited)
    assert cleaned == extract_lay_script("Here is the script.\n\n" + credited)
    assert "Kyler Rasmussen" not in cleaned
    assert "Interpreting Interpreter: Copper Tokens" in cleaned
    assert "Credit: Wikimedia Commons" in cleaned
    assert "the author’s hedge" in cleaned
    main, lay = split_lay_academic(SAMPLE_SUMMARY_MARKDOWN.rstrip() + "\n\n" + credited)
    assert "Kyler Rasmussen" not in lay
    assert "By Kyler Rasmussen" not in main
    assert main.startswith("# Interpreting Interpreter: Tokens, Not Tonnage")
    data = lay_academic_to_docx(credited)
    document = Document(BytesIO(data))
    blob = "\n".join(paragraph.text for paragraph in document.paragraphs)
    assert "Kyler Rasmussen" not in blob
    assert "Interpreting Interpreter: Copper Tokens" in blob
    assert document.core_properties.author != "Kyler Rasmussen"
    assert b"Kyler Rasmussen" not in data
    summary = Document(BytesIO(markdown_to_docx(SAMPLE_SUMMARY_MARKDOWN)))
    assert summary.core_properties.author == "Kyler Rasmussen"
    assert summary.paragraphs[0].text.lstrip().startswith("Interpreting Interpreter: Tokens, Not Tonnage")
    assert [cell.text for cell in summary.tables[0].rows[0].cells] == ["#", "Text", "Image"]


def _assert_theme_and_no_page_field(data: bytes) -> None:
    with zipfile.ZipFile(BytesIO(data)) as archive:
        theme = ET.fromstring(archive.read("word/theme/theme1.xml"))
        minor = next(theme.iter(A_NS + "minorFont"))
        latin = minor.find(A_NS + "latin")
        assert latin is not None and latin.get("typeface") == "Calibri"
        document_xml = archive.read("word/document.xml")
        root = ET.fromstring(document_xml)
        for fonts in root.iter(W_NS + "rFonts"):
            for attr in ("ascii", "hAnsi", "cs", "eastAsia"):
                assert fonts.get(W_NS + attr) in (None, "")
        assert b"PAGE" not in document_xml
        assert b"w:instrText" not in document_xml


@pytest.mark.asyncio
async def test_summarize_keeps_the_lay_script_out_of_the_main_post(tmp_path):
    pdf_path = write_sample_pdf(tmp_path / "article.pdf")
    settings = Settings(xai_api_key="test-key", xai_model="grok-4.6")
    combined = SAMPLE_SUMMARY_MARKDOWN.rstrip() + "\n\n" + LAY_SAMPLE
    seen: list[dict] = []

    def respond(request: httpx.Request) -> httpx.Response:
        if request.method == "POST" and request.url.path == "/v1/files":
            return httpx.Response(200, json={"id": "file-123"})
        if request.method == "DELETE":
            return httpx.Response(200, json={"deleted": True})
        if request.url.path != "/v1/responses":
            return httpx.Response(404)
        payload = json.loads(request.content.decode() or "{}")
        seen.append(payload)
        system = ""
        for item in payload.get("input") or []:
            if item.get("role") == "system" and isinstance(item.get("content"), str):
                system = item["content"]
        if "Output only the Lay Academic script" in system:
            text = LAY_SAMPLE
            assert payload.get("reasoning") == {"effort": "low"}
            assert "input_file" not in request.content.decode()
        elif "Return only JSON" in system:
            text = '{"findings": []}'
        else:
            text = combined
        event = json.dumps({"type": "response.output_text.delta", "delta": text})
        return httpx.Response(200, text=f"data: {event}\n\ndata: [DONE]\n\n")

    transport = httpx.MockTransport(respond)
    async with httpx.AsyncClient(transport=transport, base_url="https://api.x.ai") as raw:
        client = GrokClient(settings, client=raw)
        result = await summarize_pdf(
            pdf_path.read_bytes(),
            "70_09 (Nov26) Gee.pdf",
            settings,
            client=client,
            model_verify=False,
        )
    assert "Lay Academic" not in result.markdown
    assert "Narration (spoken)" not in result.markdown
    assert "see you next time" in result.markdown
    assert result.lay_markdown.startswith("# Copper Tokens: Lay Academic Script")
    assert result.lay_row_count == 10
    assert "on target" in result.lay_length_note
    assert len(seen) == 2
    main_doc = Document(BytesIO(markdown_to_docx(result.markdown)))
    assert [cell.text for cell in main_doc.tables[0].rows[0].cells] == ["#", "Text", "Image"]
    lay_doc = Document(BytesIO(lay_academic_to_docx(result.lay_markdown)))
    assert [cell.text for cell in lay_doc.tables[0].rows[0].cells][1:] == [
        "Narration (spoken)",
        "On-screen caption",
        "Image",
    ]


def test_cli_writes_the_lay_docx_beside_the_ai_summary(tmp_path, monkeypatch, capsys):
    pdf_path = write_sample_pdf(tmp_path / "article.pdf")

    class _Result:
        markdown = SAMPLE_SUMMARY_MARKDOWN
        lay_markdown = LAY_SAMPLE
        check_report = "# Check report\n\nNothing unresolved.\n"
        page_count = 3
        model = "grok-4.6"

        def verify_summary_line(self) -> str:
            return "ok"

    async def fake_summarize(*args, **kwargs):
        return _Result()

    monkeypatch.setattr("interpreter_summary.cli.summarize_pdf", fake_summarize)
    docx_path = tmp_path / "70_09 (Nov26) Gee.AI.docx"
    main(
        ["summarize", str(pdf_path), "--docx", str(docx_path), "--no-verify"]
    )
    lay_path = tmp_path / "70_09 (Nov26) Gee - Lay Academic Script.docx"
    assert docx_path.exists()
    assert lay_path.exists()
    main_doc = Document(docx_path)
    assert [cell.text for cell in main_doc.tables[0].rows[0].cells] == ["#", "Text", "Image"]
    assert "Lay Academic" not in "\n".join(paragraph.text for paragraph in main_doc.paragraphs)
    lay_doc = Document(lay_path)
    assert lay_doc.tables[0].rows[0].cells[1].text == "Narration (spoken)"
    assert "Post Caption" in [paragraph.text for paragraph in lay_doc.paragraphs]
    captured = capsys.readouterr()
    assert "Lay Academic Script.docx" in captured.out
    assert "XAI_API_KEY" not in captured.out
    assert "XAI_API_KEY" not in captured.err


def test_export_lay_docx_endpoint_is_separate_from_the_summary_docx():
    client = TestClient(app)
    home = client.get("/")
    assert home.status_code == 200
    assert 'id="lay-docx-btn"' in home.text
    assert "hidden" in home.text
    lay = client.post(
        "/api/export/lay-docx",
        json={"markdown": LAY_SAMPLE, "filename": "70_09 (Nov26) Gee.AI.docx"},
    )
    assert lay.status_code == 200
    assert "Gee - Lay Academic Script.docx" in lay.headers["content-disposition"]
    lay_doc = Document(BytesIO(lay.content))
    assert lay_doc.tables[0].rows[0].cells[1].text == "Narration (spoken)"
    summary = client.post("/api/export/docx", json={"markdown": SAMPLE_SUMMARY_MARKDOWN})
    assert summary.status_code == 200
    summary_doc = Document(BytesIO(summary.content))
    assert [cell.text for cell in summary_doc.tables[0].rows[0].cells] == ["#", "Text", "Image"]
    assert "Narration (spoken)" not in summary_doc.tables[0].rows[0].cells[1].text
