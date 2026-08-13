from __future__ import annotations

from pathlib import Path

from docx import Document
from docx.shared import Pt
from fpdf import FPDF

SAMPLE_ARTICLE_TITLE = "Copper, Covenants, and the Case of the Missing Ingots"
SAMPLE_AUTHOR = "A. Sample Scholar"

SAMPLE_PDF_PARAGRAPHS = [
    (
        1,
        [
            f"{SAMPLE_ARTICLE_TITLE}",
            f"by {SAMPLE_AUTHOR}",
            "Interpreter: A Journal of Latter-day Saint Faith and Scholarship 99 (2099): 1-8",
            "Abstract: This sample article exists only to exercise the summarizer. It argues that "
            "references to 'bright copper' in a fictional Nephite warehouse list are better read "
            "as covenant tokens than as a metallurgical report, because the surrounding clauses "
            "echo treaty language rather than assay notes.",
            "[Page 1]",
            "The warehouse fragment under discussion -- hereafter Fragment W -- lists 'ten ingots of "
            "bright copper' immediately after an oath formula. Critics have treated the ingots as "
            "a straightforward inventory of trade goods. That reading is possible, but it ignores "
            "how the fragment frames the metal.",
        ],
    ),
    (
        2,
        [
            "[Page 2]",
            "Three features of the surrounding clauses matter. First, the verb translated 'weigh' "
            "elsewhere in the corpus regularly means 'confirm' when the object is a promise rather "
            "than a commodity. Second, the adjective 'bright' clusters with words for holiness, not "
            "with words for ore. Third, the list ends with a witness formula ('in the presence of "
            "three') that would be odd in a shipping receipt and ordinary in a covenant text.",
            "Taken together, these features suggest the ingots function as tokens that the covenant "
            "has been 'made sure,' not as a report of what sat on a shelf.",
        ],
    ),
    (
        3,
        [
            "[Page 3]",
            "A comparative table of treaty deposits from the Late Bronze Age shows similar pairing "
            "of metal tokens with oath witnesses. The analogy is suggestive rather than decisive: "
            "Fragment W is short, and we lack the preceding column. Still, the burden of proof sits "
            "with the inventory reading, which must explain the witness formula as decorative.",
            "As I conclude: although copper can be cargo, in this fragment it more plausibly "
            "memorializes a promise. Readers who come to the passage looking for a foundry will "
            "leave disappointed; readers who come looking for a covenant will find one hiding in "
            "plain sight.",
        ],
    ),
]

SAMPLE_STYLE_SECTIONS = [
    (
        "Interpreting Interpreter: Tokens, Not Tonnage",
        "heading",
    ),
    (
        'This post is a summary of the article "Copper, Covenants, and the Case of the Missing Ingots" '
        "by A. Sample Scholar in Volume 99 of Interpreter: A Journal of Latter-day Saint Faith and "
        "Scholarship. All of the Interpreting Interpreter articles may be seen at "
        "https://interpreterfoundation.org/category/summaries/. An introduction to the Interpreting "
        "Interpreter series is available at "
        "https://interpreterfoundation.org/interpreting-interpreter-on-abstracting-thought/.",
        "body",
    ),
    ("The Takeaway", "heading"),
    (
        "Scholar argues that the 'bright copper' of Fragment W is a covenant token rather than a "
        "warehouse inventory, because the surrounding clauses use oath, witness, and holiness language "
        "instead of assay language.",
        "body",
    ),
    ("The Summary", "heading"),
    (
        "In this article, A. Sample Scholar revisits a short warehouse fragment that lists ten ingots "
        'of bright copper after an oath formula (link to "The warehouse fragment"; page 1). Critics '
        "have read the list as trade goods. Scholar instead walks through three linguistic features -- "
        "the verb often rendered 'weigh,' the adjective 'bright,' and a closing witness formula -- that "
        "fit treaty deposits better than shipping receipts (link to \"Three features\"; page 2).",
        "body",
    ),
    (
        "A comparative table of Late Bronze Age treaty deposits is offered as supporting context, with "
        "the important caveat that Fragment W is incomplete (link to \"A comparative table\"; page 3). "
        "Scholar concludes that copper can be cargo, but in this fragment it more plausibly memorializes "
        'a promise (link to "although copper can be cargo"; page 3).',
        "body",
    ),
    ("The Reflection", "heading"),
    (
        "I like how little of the case depends on discovering a new artifact and how much depends on "
        "reading the words we already have with a different set of expectations. The analogy to treaty "
        "deposits is the shakiest board in the floor, and Scholar is right to say so. Even so, once you "
        "have seen the witness formula sitting where a weight total ought to be, it is hard to go back "
        "to treating the passage as a packing list.",
        "body",
    ),
    ("Video Script", "heading"),
    (
        "[This section is included only as a style sample. The summarizer must ignore it.]",
        "body",
    ),
    (
        "You may have seen last week's video on metal imagery in Restoration scripture. This week we "
        "look at a tiny warehouse fragment and a surprisingly large claim: those ten ingots of bright "
        "copper might not be cargo at all. They might be a covenant you can hold in your hand. Check "
        "the links in the description for the written summary and the full article.",
        "body",
    ),
]


class SamplePDF(FPDF):
    def header(self) -> None:
        self.set_font("Times", "I", 9)
        self.cell(0, 8, "Interpreter: A Journal of Latter-day Saint Faith and Scholarship (sample)", align="C")
        self.ln(12)

    def footer(self) -> None:
        self.set_y(-15)
        self.set_font("Times", "I", 9)
        self.cell(0, 8, f"Page {self.page_no()}", align="C")


def write_sample_pdf(path: Path) -> Path:
    pdf = SamplePDF()
    pdf.set_auto_page_break(auto=True, margin=20)
    for _page_number, paragraphs in SAMPLE_PDF_PARAGRAPHS:
        pdf.add_page()
        for index, paragraph in enumerate(paragraphs):
            if index == 0:
                pdf.set_font("Times", "B", 16)
            elif index == 1:
                pdf.set_font("Times", "I", 12)
            else:
                pdf.set_font("Times", "", 12)
            pdf.multi_cell(0, 6, paragraph)
            pdf.ln(3)
    path.parent.mkdir(parents=True, exist_ok=True)
    pdf.output(str(path))
    return path


def write_sample_docx(path: Path) -> Path:
    document = Document()
    style = document.styles["Normal"]
    style.font.name = "Georgia"
    style.font.size = Pt(11)
    for text, kind in SAMPLE_STYLE_SECTIONS:
        if kind == "heading":
            document.add_heading(text, level=1)
        else:
            document.add_paragraph(text)
    path.parent.mkdir(parents=True, exist_ok=True)
    document.save(str(path))
    return path


def generate_samples(directory: Path) -> tuple[Path, Path]:
    pdf_path = write_sample_pdf(directory / "sample_article.pdf")
    docx_path = write_sample_docx(directory / "interpreting_interpreter_style.docx")
    return pdf_path, docx_path
