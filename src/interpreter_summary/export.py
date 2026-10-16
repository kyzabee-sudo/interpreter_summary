from __future__ import annotations

import re
from importlib.resources import as_file, files
from io import BytesIO

from docx import Document
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Inches, Pt, RGBColor, Twips

HEADING_RE = re.compile(r"^(#{1,3})\s+(.*)$")
HEADING_AT_LINE_START_RE = re.compile(r"(?m)^#{1,3}[ \t]+\S")
HEADING_ANYWHERE_RE = re.compile(r"#{1,3}[ \t]+\S")
# The series title is the start of the post even when chatter is glued to it
# on the same line. A later "## The Takeaway" must not win that race.
TITLE_HEADING_RE = re.compile(r"# Interpreting Interpreter\b", re.IGNORECASE)
REFLECTION_RE = re.compile(r"^#{1,3}\s+(?:the\s+)?reflection\s*:?\s*$", re.IGNORECASE)
RESUME_HEADING_RE = re.compile(r"^#{1,2}\s+\S")
SEPARATOR_CELL_RE = re.compile(r":?-{3,}:?")
# **markers** are stripped. *italics* stay. Locators are coloured, never bold.
BOLD_MARKER_RE = re.compile(r"\*\*([^*]+)\*\*")
ITALIC_RE = re.compile(r"\*([^*]+)\*")
LOCATOR_SPAN_RE = re.compile(
    r"(?:\[[^\]]+\]\s*)?\(link to\s+[“\"”][^“\"”]+[“\"”];\s*page\s+\d+\)",
    re.IGNORECASE,
)
LINK_QUOTE_RE = re.compile(
    r"(\(link to\s+)[“\"”]([^“\"”]+)[“\"”](\s*;\s*page\s+\d+\))",
    re.IGNORECASE,
)
LOCATOR_RED = RGBColor(0xFF, 0, 0)

# Column widths from Kyler's template (twips): #, Text, Image. Sum is tblW 9001.
VIDEO_WIDTHS = (450, 6325, 2226)


def section_key(title: str) -> str | None:
    lowered = title.strip().rstrip(":").lower()
    lowered = re.sub(r"\s+", " ", lowered)
    if lowered in {"the takeaway", "takeaway"}:
        return "takeaway"
    if lowered in {"the q&a", "q&a", "the q & a", "q & a", "questions and answers", "the questions and answers"}:
        return "qa"
    if lowered in {"the summary", "summary"}:
        return "summary"
    if lowered in {"video script", "the video script", "script"}:
        return "video_script"
    if lowered in {"the reflection", "reflection"}:
        return "reflection"
    return None


def parse_sections(markdown: str) -> dict[str, str]:
    """Split a generated draft into named sections for the UI and length checks."""
    sections = {
        "title": "",
        "intro": "",
        "takeaway": "",
        "qa": "",
        "summary": "",
        "video_script": "",
    }
    current = "intro"
    buckets: dict[str, list[str]] = {key: [] for key in sections if key != "title"}
    for raw in markdown.replace("\r\n", "\n").split("\n"):
        heading = HEADING_RE.match(raw.strip())
        if heading and len(heading.group(1)) <= 2:
            title = heading.group(2).strip()
            if len(heading.group(1)) == 1:
                sections["title"] = title
                current = "intro"
                continue
            key = section_key(title)
            if key:
                current = key
                continue
        if current == "reflection":
            continue
        buckets[current].append(raw)
    for key, lines in buckets.items():
        sections[key] = "\n".join(lines).strip()
    return sections


def strip_preamble(markdown: str) -> str:
    """Drop model chatter before the post, including chatter glued to the title line.

    The series title wins even when it shares a line with the chatter, and even
    when that line also contains a false heading such as "page # 1". A later
    "## The Takeaway" must not be treated as the start of the post.
    """
    text = markdown.replace("\r\n", "\n").lstrip("\ufeff")
    title = TITLE_HEADING_RE.search(text)
    if title is not None:
        start = title.start()
    else:
        match = HEADING_AT_LINE_START_RE.search(text) or HEADING_ANYWHERE_RE.search(text)
        start = match.start() if match else 0
    if start > 0:
        text = text[start:]
    stripped = text.strip()
    return f"{stripped}\n" if stripped else ""


def strip_reflection_section(markdown: str) -> str:
    """Drop a Reflection block if the model still emits the retired section."""
    lines = markdown.replace("\r\n", "\n").split("\n")
    kept: list[str] = []
    skipping = False
    for line in lines:
        stripped = line.strip()
        if REFLECTION_RE.match(stripped):
            skipping = True
            continue
        if skipping and RESUME_HEADING_RE.match(stripped) and not REFLECTION_RE.match(stripped):
            skipping = False
        if skipping:
            continue
        kept.append(line)
    text = re.sub(r"\n{3,}", "\n\n", "\n".join(kept)).strip()
    return f"{text}\n" if text else ""


def markdown_to_docx(markdown: str) -> bytes:
    """Write a draft that follows Kyler's Word template.

    The template supplies Calibri through the theme minor font, docDefaults
    (11 pt, 8 pt after, 1.08 line spacing), Table Grid, and List Paragraph
    numbering. Nothing here sets a font name or single line spacing.
    """
    document = _document_from_template()
    _apply_page_setup(document)
    document.core_properties.author = "Kyler Rasmussen"

    lines = markdown.replace("\r\n", "\n").split("\n")
    index = 0
    while index < len(lines):
        line = lines[index]
        if not line.strip():
            index += 1
            continue
        if _is_table_start(lines, index):
            index = _add_table(document, lines, index)
            continue
        heading = HEADING_RE.match(line.strip())
        if heading:
            level = len(heading.group(1))
            title = heading.group(2).strip()
            if level == 1:
                _add_title(document, title)
            elif level == 2:
                _add_section_heading(document, title)
            else:
                _add_question(document, title)
            index += 1
            continue
        if line.startswith(">"):
            parts: list[str] = []
            while index < len(lines) and lines[index].startswith(">"):
                body = lines[index][2:] if lines[index].startswith("> ") else lines[index][1:]
                parts.append(body.strip())
                index += 1
            _add_quote(document, " ".join(part for part in parts if part))
            continue
        if line.startswith(("- ", "* ")):
            _add_bullet(document, line[2:].strip())
            index += 1
            continue
        if re.fullmatch(r"(-{3,}|\*{3,})", line.strip()):
            index += 1
            continue
        _add_body(document, line.strip())
        index += 1

    buffer = BytesIO()
    document.save(buffer)
    return buffer.getvalue()


def _document_from_template() -> Document:
    """Open Kyler's template and drop its placeholder body, keeping sectPr."""
    resource = files("interpreter_summary").joinpath("style_assets", "summary_template.docx")
    with as_file(resource) as path:
        document = Document(str(path))
    body = document.element.body
    for child in list(body):
        if child.tag == qn("w:sectPr"):
            continue
        body.remove(child)
    return document


def _apply_page_setup(document: Document) -> None:
    section = document.sections[0]
    section.page_width = Inches(8.5)
    section.page_height = Inches(11)
    section.left_margin = Inches(1)
    section.right_margin = Inches(1)
    section.top_margin = Inches(1)
    section.bottom_margin = Inches(1)


def _add_title(document: Document, text: str) -> None:
    paragraph = document.add_paragraph()
    paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER
    match = re.match(r"^Interpreting Interpreter:\s*(.*)$", text)
    _add_run(paragraph, " ", bold=True, size=16)
    if match:
        _add_run(paragraph, "Interpreting ", bold=True, size=16)
        _add_run(paragraph, "Interpreter", bold=True, italic=True, size=16)
        _add_run(paragraph, f": {match.group(1)}", bold=True, size=16)
    else:
        _add_run(paragraph, text, bold=True, size=16)
    blank = document.add_paragraph()
    blank.alignment = WD_ALIGN_PARAGRAPH.CENTER


def _add_section_heading(document: Document, text: str) -> None:
    paragraph = document.add_paragraph()
    if section_key(text) == "video_script":
        _add_run(paragraph, text, bold=False, italic=True)
    else:
        _add_run(paragraph, text, bold=True, italic=False)


def _add_question(document: Document, text: str) -> None:
    paragraph = document.add_paragraph()
    _add_inline(paragraph, text, bold=True, italic=True)


def _add_body(document: Document, text: str) -> None:
    paragraph = document.add_paragraph()
    _add_inline(paragraph, text, bold=False, italic=False)


def _add_quote(document: Document, text: str) -> None:
    paragraph = document.add_paragraph()
    paragraph.paragraph_format.left_indent = Twips(720)
    _set_spacing(paragraph, after="0")
    _add_inline(paragraph, text, bold=False, italic=False)


def _add_bullet(document: Document, text: str) -> None:
    paragraph = document.add_paragraph(style="List Paragraph")
    _apply_bullet_numbering(paragraph)
    _add_inline(paragraph, text, bold=False, italic=False)


def _apply_bullet_numbering(paragraph) -> None:
    """List Paragraph plus Kyler's '-' bullet (numId 1) and contextual spacing."""
    p_pr = paragraph._p.get_or_add_pPr()
    num_pr = OxmlElement("w:numPr")
    ilvl = OxmlElement("w:ilvl")
    ilvl.set(qn("w:val"), "0")
    num_id = OxmlElement("w:numId")
    num_id.set(qn("w:val"), "1")
    num_pr.append(ilvl)
    num_pr.append(num_id)
    p_style = p_pr.find(qn("w:pStyle"))
    if p_style is not None:
        p_style.addnext(num_pr)
    else:
        p_pr.insert(0, num_pr)
    if p_pr.find(qn("w:contextualSpacing")) is None:
        p_pr.append(OxmlElement("w:contextualSpacing"))


def _set_spacing(paragraph, *, after: str) -> None:
    """Set space-after without dropping the template's 1.08 line spacing."""
    p_pr = paragraph._p.get_or_add_pPr()
    spacing = p_pr.find(qn("w:spacing"))
    if spacing is None:
        spacing = OxmlElement("w:spacing")
        p_pr.append(spacing)
    spacing.set(qn("w:after"), after)
    spacing.set(qn("w:line"), "259")
    spacing.set(qn("w:lineRule"), "auto")


def _prepare_inline(text: str) -> str:
    stripped = BOLD_MARKER_RE.sub(r"\1", text)
    return LINK_QUOTE_RE.sub(
        lambda match: f"{match.group(1)}\u201c{match.group(2)}\u201d{match.group(3)}",
        stripped,
    )


def _inline_spans(prepared: str) -> list[tuple[int, int, str]]:
    spans: list[tuple[int, int, str]] = []
    for match in LOCATOR_SPAN_RE.finditer(prepared):
        spans.append((match.start(), match.end(), "locator"))
    for match in ITALIC_RE.finditer(prepared):
        start, end = match.start(), match.end()
        if any(start < stop and end > begin for begin, stop, _kind in spans):
            continue
        spans.append((start, end, "italic"))
    spans.sort()
    return spans


def _add_inline(paragraph, text: str, *, bold: bool, italic: bool) -> None:
    """Emit plain text. **markers** are dropped. Locators are red and never bold."""
    prepared = _prepare_inline(text)
    cursor = 0
    for start, end, kind in _inline_spans(prepared):
        if start < cursor:
            continue
        if start > cursor:
            _add_run(paragraph, prepared[cursor:start], bold=bold, italic=italic)
        if kind == "locator":
            _add_run(paragraph, prepared[start:end], bold=False, italic=False, color=LOCATOR_RED)
        else:
            _add_run(paragraph, prepared[start + 1 : end - 1], bold=bold, italic=True)
        cursor = end
    if cursor < len(prepared):
        _add_run(paragraph, prepared[cursor:], bold=bold, italic=italic)
    if not prepared:
        _add_run(paragraph, "", bold=bold, italic=italic)


def _add_run(
    paragraph,
    text: str,
    *,
    bold: bool,
    italic: bool = False,
    size: int | None = None,
    color: RGBColor | None = None,
):
    run = paragraph.add_run(text)
    run.bold = bold
    run.italic = italic
    if size is not None:
        run.font.size = Pt(size)
    if color is not None:
        run.font.color.rgb = color
    return run


def _is_table_start(lines: list[str], index: int) -> bool:
    if not lines[index].lstrip().startswith("|"):
        return False
    return index + 1 < len(lines) and lines[index + 1].lstrip().startswith("|")


def _split_row(line: str) -> list[str]:
    stripped = line.strip()
    if stripped.startswith("|"):
        stripped = stripped[1:]
    if stripped.endswith("|"):
        stripped = stripped[:-1]
    return [cell.strip() for cell in stripped.split("|")]


def _is_separator(cells: list[str]) -> bool:
    return bool(cells) and all(SEPARATOR_CELL_RE.fullmatch(cell.replace(" ", "")) for cell in cells)


def _add_table(document: Document, lines: list[str], index: int) -> int:
    rows: list[list[str]] = []
    while index < len(lines) and lines[index].lstrip().startswith("|"):
        cells = _split_row(lines[index])
        if not _is_separator(cells):
            rows.append(cells)
        index += 1
    if not rows:
        return index
    columns = max(len(row) for row in rows)
    table = document.add_table(rows=len(rows), cols=columns)
    table.style = "Table Grid"
    widths = _column_widths(columns)
    _set_table_widths(table, widths)
    for row_index, row in enumerate(rows):
        for col_index in range(columns):
            value = row[col_index] if col_index < len(row) else ""
            cell = table.rows[row_index].cells[col_index]
            paragraph = cell.paragraphs[0]
            if row_index == 0:
                paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER
            _add_inline(paragraph, value, bold=row_index == 0, italic=False)
            _set_cell_width(cell, widths[col_index])
    return index


def _column_widths(columns: int) -> tuple[int, ...]:
    if columns == 3:
        return VIDEO_WIDTHS
    total = sum(VIDEO_WIDTHS)
    each = total // max(columns, 1)
    return tuple(each for _ in range(columns))


def _set_table_widths(table, widths: tuple[int, ...]) -> None:
    tbl = table._tbl
    tbl_pr = tbl.tblPr
    layout = tbl_pr.find(qn("w:tblLayout"))
    if layout is not None:
        tbl_pr.remove(layout)
    tbl_w = tbl_pr.find(qn("w:tblW"))
    if tbl_w is None:
        tbl_w = OxmlElement("w:tblW")
        tbl_pr.append(tbl_w)
    total = sum(widths)
    tbl_w.set(qn("w:w"), str(total))
    tbl_w.set(qn("w:type"), "dxa")
    grid = tbl.find(qn("w:tblGrid"))
    if grid is None:
        grid = OxmlElement("w:tblGrid")
        tbl_pr.addnext(grid)
    else:
        for child in list(grid):
            grid.remove(child)
    for width in widths:
        column = OxmlElement("w:gridCol")
        column.set(qn("w:w"), str(width))
        grid.append(column)


def _set_cell_width(cell, width: int) -> None:
    tc_pr = cell._tc.get_or_add_tcPr()
    tc_w = tc_pr.find(qn("w:tcW"))
    if tc_w is None:
        tc_w = OxmlElement("w:tcW")
        tc_pr.append(tc_w)
    tc_w.set(qn("w:w"), str(width))
    tc_w.set(qn("w:type"), "dxa")
    cell.width = Twips(width)
