from __future__ import annotations

from pathlib import Path

from docx import Document
from docx.oxml.ns import qn

PACKAGE_STYLE = Path(__file__).resolve().parent / "style_assets" / "style_guide.md"


def load_style_text(path: Path | None = None) -> str:
    """Load a Markdown or Word style document, including any video-script table."""
    target = path or PACKAGE_STYLE
    if not target.exists():
        raise FileNotFoundError(f"Style file not found: {target}")
    suffix = target.suffix.lower()
    if suffix in {".md", ".txt"}:
        text = target.read_text(encoding="utf-8")
    elif suffix in {".docx"}:
        text = _docx_to_text(target)
    else:
        raise ValueError(f"Unsupported style file type: {suffix}")
    return text.strip() + "\n"


def load_style_bytes(data: bytes, filename: str) -> str:
    suffix = Path(filename).suffix.lower()
    if suffix in {".md", ".txt"}:
        text = data.decode("utf-8")
    elif suffix == ".docx":
        from io import BytesIO

        text = _docx_to_text(BytesIO(data))
    else:
        raise ValueError(f"Unsupported style file type: {suffix}")
    return text.strip() + "\n"


def _docx_to_text(source) -> str:
    document = Document(source)
    chunks: list[str] = []
    paragraphs = document.paragraphs
    tables = document.tables
    paragraph_index = 0
    table_index = 0
    for child in document.element.body.iterchildren():
        if child.tag == qn("w:p"):
            chunks.append(paragraphs[paragraph_index].text)
            paragraph_index += 1
        elif child.tag == qn("w:tbl"):
            chunks.append(_table_to_markdown(tables[table_index]))
            table_index += 1
    return "\n".join(chunks).strip() + "\n"


def _table_to_markdown(table) -> str:
    rows: list[str] = []
    for row in table.rows:
        cells = [cell.text.replace("\n", " ").replace("|", "/").strip() for cell in row.cells]
        rows.append("| " + " | ".join(cells) + " |")
    if not rows:
        return ""
    separator = "| " + " | ".join(["---"] * len(table.rows[0].cells)) + " |"
    return "\n".join([rows[0], separator, *rows[1:]])
